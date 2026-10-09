import openpyxl
import glob
import json
import os
from datetime import datetime
from collections import defaultdict

def main():
    files = glob.glob('/Users/Admin/Documents/KIng BLue/**/So_chi_tiet_ban_hang 2909.xlsx', recursive=True)
    if not files:
        print("Error: Excel file not found!")
        return
    filepath = files[0]
    print(f"Loading data from: {filepath}")

    wb = openpyxl.load_workbook(filepath, read_only=True)

    customers = {}
    monthly_stats = defaultdict(lambda: {'gross': 0, 'discount': 0, 'return': 0, 'net': 0, 'orders': set(), 'customers': set(), 'items': 0})
    province_stats = defaultdict(lambda: {'gross': 0, 'discount': 0, 'return': 0, 'net': 0, 'orders': set(), 'customers': set()})
    sales_rep_stats = defaultdict(lambda: {'gross': 0, 'discount': 0, 'return': 0, 'net': 0, 'orders': set(), 'customers': set()})
    category_stats = defaultdict(lambda: {'gross': 0, 'qty': 0, 'customers': set(), 'orders': set()})
    product_stats = defaultdict(lambda: {'code': '', 'name': '', 'category': '', 'gross': 0, 'qty': 0, 'customers': set()})

    ref_date = datetime(2026, 9, 30)

    total_gross = 0.0
    total_discount = 0.0
    total_return = 0.0
    total_net = 0.0
    all_orders = set()
    total_line_items = 0
    total_qty = 0.0

    sheet_names = ['SỔ CHI TIẾT BÁN HÀNG', 'SỔ CHI TIẾT BÁN HÀNG_1']

    for sname in sheet_names:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        print(f"Reading sheet: {sname}...")
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r or r[0] is None or r[3] is None:
                continue
            
            total_line_items += 1
            raw_date = r[0]
            if isinstance(raw_date, datetime):
                dt = raw_date
            else:
                try:
                    dt = datetime.strptime(str(raw_date)[:10], '%Y-%m-%d')
                except:
                    dt = datetime(2026, 1, 1)
            
            month_key = dt.strftime('%Y-%m')
            so_ct = str(r[1]).strip() if r[1] is not None else ''
            ma_kh = str(r[3]).strip()
            ten_kh = str(r[4]).strip() if r[4] is not None else ma_kh
            dia_chi = str(r[5]).strip() if r[5] is not None else ''
            sdt = str(r[6]).strip() if r[6] is not None else ''
            ma_hang = str(r[7]).strip() if r[7] is not None else ''
            ten_hang = str(r[8]).strip() if r[8] is not None else ma_hang
            sl_ban = float(r[9] or 0)
            doanh_so = float(r[19] or 0)
            chiet_khau = float(r[24] or 0)
            tra_lai = float(r[32] or 0)
            nvbh = str(r[38]).strip() if r[38] is not None else 'Chưa gán'
            tinh = str(r[39]).strip() if r[39] is not None else 'Chưa xác định'
            nhom = str(r[42]).strip() if r[42] is not None else 'Khác'
            
            if not tinh or tinh == 'None': tinh = 'Chưa xác định'
            if not nvbh or nvbh == 'None': nvbh = 'Chưa gán'
            if not nhom or nhom == 'None': nhom = 'Khác'

            net_rev = doanh_so - chiet_khau - tra_lai
            
            total_gross += doanh_so
            total_discount += chiet_khau
            total_return += tra_lai
            total_net += net_rev
            total_qty += sl_ban
            if so_ct: all_orders.add(so_ct)
                
            m = monthly_stats[month_key]
            m['gross'] += doanh_so
            m['discount'] += chiet_khau
            m['return'] += tra_lai
            m['net'] += net_rev
            m['items'] += sl_ban
            if so_ct: m['orders'].add(so_ct)
            m['customers'].add(ma_kh)
            
            p = province_stats[tinh]
            p['gross'] += doanh_so
            p['discount'] += chiet_khau
            p['return'] += tra_lai
            p['net'] += net_rev
            if so_ct: p['orders'].add(so_ct)
            p['customers'].add(ma_kh)
            
            s = sales_rep_stats[nvbh]
            s['gross'] += doanh_so
            s['discount'] += chiet_khau
            s['return'] += tra_lai
            s['net'] += net_rev
            if so_ct: s['orders'].add(so_ct)
            s['customers'].add(ma_kh)
            
            c = category_stats[nhom]
            c['gross'] += doanh_so
            c['qty'] += sl_ban
            c['customers'].add(ma_kh)
            if so_ct: c['orders'].add(so_ct)
            
            pr = product_stats[ma_hang]
            pr['code'] = ma_hang
            pr['name'] = ten_hang
            pr['category'] = nhom
            pr['gross'] += doanh_so
            pr['qty'] += sl_ban
            pr['customers'].add(ma_kh)
            
            if ma_kh not in customers:
                customers[ma_kh] = {
                    'id': ma_kh,
                    'name': ten_kh,
                    'address': dia_chi,
                    'phone': sdt,
                    'province': tinh,
                    'sales_reps': defaultdict(int),
                    'first_date': dt,
                    'last_date': dt,
                    'orders': set(),
                    'line_items': 0,
                    'total_qty': 0,
                    'gross': 0.0,
                    'discount': 0.0,
                    'return': 0.0,
                    'net': 0.0,
                    'categories': defaultdict(float),
                    'top_products': defaultdict(lambda: {'name': '', 'gross': 0, 'qty': 0})
                }
            
            cust = customers[ma_kh]
            if dt < cust['first_date']: cust['first_date'] = dt
            if dt > cust['last_date']: cust['last_date'] = dt
            if so_ct: cust['orders'].add(so_ct)
            cust['line_items'] += 1
            cust['total_qty'] += sl_ban
            cust['gross'] += doanh_so
            cust['discount'] += chiet_khau
            cust['return'] += tra_lai
            cust['net'] += net_rev
            cust['sales_reps'][nvbh] += 1
            cust['categories'][nhom] += doanh_so
            
            prod_entry = cust['top_products'][ma_hang]
            prod_entry['name'] = ten_hang
            prod_entry['gross'] += doanh_so
            prod_entry['qty'] += sl_ban

    sorted_customers = sorted(customers.values(), key=lambda x: x['net'], reverse=True)

    # Pareto & ABC Analysis
    cum_rev = 0.0
    for c in sorted_customers:
        cum_rev += c['net']
        c_pct = cum_rev / total_net if total_net > 0 else 0
        if c_pct <= 0.80:
            c['abc'] = 'A'
        elif c_pct <= 0.95:
            c['abc'] = 'B'
        else:
            c['abc'] = 'C'

    # RFM Scoring & Segmentation
    for c in sorted_customers:
        r_days = (ref_date - c['last_date']).days
        f_orders = len(c['orders'])
        m_val = c['net']
        c['r_days'] = r_days
        c['f_orders'] = f_orders
        
        # Primary sales rep
        c['main_sales_rep'] = max(c['sales_reps'].items(), key=lambda x: x[1])[0] if c['sales_reps'] else 'Chưa gán'

        # Segmentation
        if r_days <= 30 and (f_orders >= 15 or m_val >= 300_000_000):
            segment = 'VIP / Khách hàng Vàng'
            segment_code = 'CHAMPION'
            segment_color = '#3b82f6'
            action = 'Chăm sóc đặc biệt, ký hợp đồng năm, tặng quà tri ân VIP, ưu tiên giữ hàng.'
        elif r_days <= 45 and (f_orders >= 8 or m_val >= 100_000_000):
            segment = 'Khách hàng Trung thành'
            segment_code = 'LOYAL'
            segment_color = '#10b981'
            action = 'Upsell các dòng máy mới, duy trì tần suất gọi điện định kỳ 1 tuần/lần.'
        elif r_days <= 45 and (f_orders >= 2 or m_val >= 30_000_000):
            segment = 'Tiềm năng phát triển'
            segment_code = 'POTENTIAL'
            segment_color = '#06b6d4'
            action = 'Đề xuất chính sách chiết khấu bậc thang theo doanh số tháng để kích cầu.'
        elif r_days <= 45:
            segment = 'Khách hàng Mới / Mua gần đây'
            segment_code = 'NEW_ACTIVE'
            segment_color = '#8b5cf6'
            action = 'Theo dõi trải nghiệm sử dụng đợt đầu, gửi catalogue giới thiệu thêm mã hàng.'
        elif 45 < r_days <= 90 and (f_orders >= 5 or m_val >= 50_000_000):
            segment = 'Cần chăm sóc / Nguy cơ nguội'
            segment_code = 'NEEDS_ATTENTION'
            segment_color = '#f59e0b'
            action = 'Sales phụ trách cần liên hệ hỏi thăm tồn kho, giới thiệu mã hàng khuyến mãi.'
        elif 90 < r_days <= 150:
            segment = 'Nguy cơ rời bỏ cao'
            segment_code = 'AT_RISK'
            segment_color = '#ef4444'
            action = 'BÁO ĐỘNG: Gặp trực tiếp chủ đại lý, tìm hiểu lý do ngưng nhập hàng, chào chính sách đặc biệt.'
        else:
            segment = 'Ngủ đông / Rời bỏ'
            segment_code = 'LOST'
            segment_color = '#64748b'
            action = 'Chiến dịch re-engagement: gửi thông báo sản phẩm mới kèm voucher ưu đãi mở lại tài khoản.'
        
        c['segment'] = segment
        c['segment_code'] = segment_code
        c['segment_color'] = segment_color
        c['action'] = action

    # Serialize customer records
    formatted_customers = []
    for c in sorted_customers:
        # Top 5 categories
        top_cats = sorted(c['categories'].items(), key=lambda x: x[1], reverse=True)[:5]
        top_prods = sorted(c['top_products'].values(), key=lambda x: x['gross'], reverse=True)[:5]
        
        formatted_customers.append({
            'id': c['id'],
            'name': c['name'],
            'address': c['address'],
            'phone': c['phone'],
            'province': c['province'],
            'main_sales_rep': c['main_sales_rep'],
            'first_date': c['first_date'].strftime('%Y-%m-%d'),
            'last_date': c['last_date'].strftime('%Y-%m-%d'),
            'r_days': c['r_days'],
            'f_orders': c['f_orders'],
            'line_items': c['line_items'],
            'total_qty': round(c['total_qty'], 1),
            'gross': round(c['gross'], 0),
            'discount': round(c['discount'], 0),
            'discount_pct': round((c['discount'] / c['gross'] * 100) if c['gross'] > 0 else 0, 2),
            'return': round(c['return'], 0),
            'net': round(c['net'], 0),
            'abc': c['abc'],
            'segment': c['segment'],
            'segment_code': c['segment_code'],
            'segment_color': c['segment_color'],
            'action': c['action'],
            'top_categories': [{'name': cat, 'gross': round(val, 0)} for cat, val in top_cats],
            'top_products': [{'name': pr['name'], 'gross': round(pr['gross'], 0), 'qty': pr['qty']} for pr in top_prods]
        })

    # Summary KPIs
    kpi_summary = {
        'total_gross': round(total_gross, 0),
        'total_discount': round(total_discount, 0),
        'discount_pct': round((total_discount / total_gross * 100) if total_gross > 0 else 0, 2),
        'total_return': round(total_return, 0),
        'return_pct': round((total_return / total_gross * 100) if total_gross > 0 else 0, 2),
        'total_net': round(total_net, 0),
        'total_orders': len(all_orders),
        'total_customers': len(customers),
        'total_line_items': total_line_items,
        'total_qty': round(total_qty, 0),
        'aov': round(total_net / len(all_orders) if all_orders else 0, 0),
        'arpu': round(total_net / len(customers) if customers else 0, 0),
        'avg_orders_per_cust': round(len(all_orders) / len(customers) if customers else 0, 1),
        'date_range': '01/01/2026 - 29/09/2026'
    }

    # Monthly Trend
    monthly_data = []
    for m in sorted(monthly_stats.keys()):
        v = monthly_stats[m]
        monthly_data.append({
            'month': m,
            'month_label': f"Tháng {int(m.split('-')[1])}",
            'gross': round(v['gross'], 0),
            'discount': round(v['discount'], 0),
            'discount_pct': round((v['discount'] / v['gross'] * 100) if v['gross'] > 0 else 0, 1),
            'return': round(v['return'], 0),
            'net': round(v['net'], 0),
            'orders': len(v['orders']),
            'customers': len(v['customers']),
            'items': round(v['items'], 0)
        })

    # Province Breakdown
    province_data = []
    sorted_prov = sorted(province_stats.items(), key=lambda x: x[1]['net'], reverse=True)
    for p, v in sorted_prov:
        province_data.append({
            'province': p,
            'net': round(v['net'], 0),
            'gross': round(v['gross'], 0),
            'discount': round(v['discount'], 0),
            'return': round(v['return'], 0),
            'customers': len(v['customers']),
            'orders': len(v['orders']),
            'pct_of_total': round((v['net'] / total_net * 100) if total_net > 0 else 0, 2)
        })

    # Sales Rep Breakdown
    sales_rep_data = []
    sorted_sales = sorted(sales_rep_stats.items(), key=lambda x: x[1]['net'], reverse=True)
    for s, v in sorted_sales:
        sales_rep_data.append({
            'name': s,
            'net': round(v['net'], 0),
            'gross': round(v['gross'], 0),
            'discount': round(v['discount'], 0),
            'customers': len(v['customers']),
            'orders': len(v['orders']),
            'arpu': round((v['net'] / len(v['customers'])) if v['customers'] else 0, 0),
            'pct_of_total': round((v['net'] / total_net * 100) if total_net > 0 else 0, 2)
        })

    # Category Breakdown
    category_data = []
    sorted_cat = sorted(category_stats.items(), key=lambda x: x[1]['gross'], reverse=True)
    for c, v in sorted_cat:
        category_data.append({
            'name': c,
            'gross': round(v['gross'], 0),
            'qty': round(v['qty'], 0),
            'customers': len(v['customers']),
            'orders': len(v['orders']),
            'pct_of_total': round((v['gross'] / total_gross * 100) if total_gross > 0 else 0, 2)
        })

    # Top Products
    product_data = []
    sorted_prod = sorted(product_stats.values(), key=lambda x: x['gross'], reverse=True)[:30]
    for p in sorted_prod:
        product_data.append({
            'code': p['code'],
            'name': p['name'],
            'category': p['category'],
            'gross': round(p['gross'], 0),
            'qty': round(p['qty'], 0),
            'customers': len(p['customers'])
        })

    # Segment Summary
    segment_summary = defaultdict(lambda: {'count': 0, 'net': 0, 'gross': 0, 'orders': 0, 'customers': []})
    for c in sorted_customers:
        seg = c['segment']
        segment_summary[seg]['count'] += 1
        segment_summary[seg]['net'] += c['net']
        segment_summary[seg]['gross'] += c['gross']
        segment_summary[seg]['orders'] += len(c['orders'])
        segment_summary[seg]['color'] = c['segment_color']
        segment_summary[seg]['code'] = c['segment_code']
        segment_summary[seg]['action'] = c['action']
        if len(segment_summary[seg]['customers']) < 5:
            segment_summary[seg]['customers'].append(c['name'])

    segment_data = []
    for seg, v in sorted(segment_summary.items(), key=lambda x: x[1]['net'], reverse=True):
        segment_data.append({
            'name': seg,
            'code': v['code'],
            'color': v['color'],
            'action': v['action'],
            'count': v['count'],
            'net': round(v['net'], 0),
            'gross': round(v['gross'], 0),
            'orders': v['orders'],
            'cust_pct': round((v['count'] / len(customers) * 100), 1),
            'rev_pct': round((v['net'] / total_net * 100), 1),
            'sample_customers': v['customers']
        })

    # ABC Summary
    abc_summary = {
        'A': {'count': 0, 'net': 0, 'label': 'Hạng A (Top 80% Doanh thu - Cốt lõi)', 'color': '#2563eb'},
        'B': {'count': 0, 'net': 0, 'label': 'Hạng B (15% Tiếp theo - Trọng tâm)', 'color': '#0ea5e9'},
        'C': {'count': 0, 'net': 0, 'label': 'Hạng C (5% Cuối cùng - Khách lẻ)', 'color': '#94a3b8'}
    }
    for c in sorted_customers:
        abc_summary[c['abc']]['count'] += 1
        abc_summary[c['abc']]['net'] += c['net']

    abc_data = []
    for tier, v in abc_summary.items():
        abc_data.append({
            'tier': tier,
            'label': v['label'],
            'color': v['color'],
            'count': v['count'],
            'net': round(v['net'], 0),
            'cust_pct': round((v['count'] / len(customers) * 100), 1),
            'rev_pct': round((v['net'] / total_net * 100), 1)
        })

    final_payload = {
        'kpi': kpi_summary,
        'monthly': monthly_data,
        'abc': abc_data,
        'segments': segment_data,
        'provinces': province_data,
        'sales_reps': sales_rep_data,
        'categories': category_data,
        'products': product_data,
        'customers': formatted_customers
    }

    out_file = '/Users/Admin/Documents/KIng BLue/Report/customer_analytics_data.json'
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(final_payload, f, ensure_ascii=False, indent=2)

    print(f"Successfully exported data to: {out_file}")
    print(f"Total customers: {len(formatted_customers)}")
    print(f"File size: {os.path.getsize(out_file) / 1024:.1f} KB")

if __name__ == '__main__':
    main()
