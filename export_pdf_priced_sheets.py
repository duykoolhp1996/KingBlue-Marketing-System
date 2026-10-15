import json
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Đọc dữ liệu đã khớp 100% từ PDF
with open("database/bang_gia_kingblue_pdf_goc.json", "r", encoding="utf-8") as f:
    pdf_catalog = json.load(f)

# Import matching logic
from match_pdf_with_sheet import matched_results

# Tạo workbook Excel
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Color Palette (King Blue Theme)
c_primary = "1B365D"  # Dark Blue
c_accent = "2E75B6"   # Medium Blue
c_gold = "C55A11"     # Highlight Gold/Orange
c_sky = "D9E1F2"
c_gray = "595959"
c_border = "D9D9D9"

font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
font_h1 = Font(name="Calibri", size=12, bold=True, color=c_primary)
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_regular = Font(name="Calibri", size=11)
font_bold = Font(name="Calibri", size=11, bold=True)
font_sub = Font(name="Calibri", size=9, italic=True, color="555555")

fill_primary = PatternFill(start_color=c_primary, end_color=c_primary, fill_type="solid")
fill_accent = PatternFill(start_color=c_accent, end_color=c_accent, fill_type="solid")
fill_gold = PatternFill(start_color="D66011", end_color="D66011", fill_type="solid")
fill_sky = PatternFill(start_color=c_sky, end_color=c_sky, fill_type="solid")
fill_header_gray = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=c_border),
    right=Side(style='thin', color=c_border),
    top=Side(style='thin', color=c_border),
    bottom=Side(style='thin', color=c_border)
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# ==============================================================================
# TAB 1: BẢNG THEO DÕI GIÁ BÁN SÀN TMĐT (1,055 DÒNG SẢN PHẨM KHỚP 100% PDF BÁO GIÁ)
# ==============================================================================
ws1 = wb.create_sheet(title="1_Theo_Doi_Gia_San_TMDT")
ws1.views.sheetView[0].showGridLines = True

headers1 = [
    "STT", "Mã Đại Lý", "Tên Đại Lý / Khách Hàng", "Tỉnh Thành", "Nhân Viên Phụ Trách", "Sàn TMĐT",
    "Mã Sản Phẩm", "Tên Sản Phẩm KingBlue", "ĐVT",
    "GIÁ SỈ (VNĐ)", "GIÁ LẺ NIÊM YẾT (VNĐ)", "GIÁ SÀN TMĐT QUY ĐỊNH (VNĐ)",
    "GIÁ BÁN THỰC TẾ TRÊN SÀN (VNĐ)", "Chênh Lệch Với Giá Sàn (VNĐ)", "% Lệch So Với Giá Sàn",
    "KẾT LUẬN TUÂN THỦ GIÁ SÀN", "Ngày Kiểm Tra", "Link Sản Phẩm / Bằng Chứng", "Hướng Xử Lý / Ghi Chú"
]

ws1.append(headers1)
for col_num in range(1, len(headers1) + 1):
    c = ws1.cell(row=1, column=col_num)
    c.font = font_header
    c.fill = fill_primary
    if col_num == 12: # Cột Giá Sàn TMĐT Quy Định
        c.fill = fill_gold
    c.alignment = align_center
    c.border = thin_border

for idx, r in enumerate(matched_results, 1):
    curr_r = idx + 1
    # Formulas:
    # Cột L: Giá Sàn TMĐT Quy Định
    # Cột M: Giá Bán Thực Tế Trên Sàn (Người dùng điền vào)
    # Cột N: Chênh Lệch Tiền = =IF(ISNUMBER(M{row}), M{row}-L{row}, "")
    # Cột O: % Chênh Lệch = =IF(AND(ISNUMBER(M{row}),ISNUMBER(L{row}),L{row}>0), (M{row}-L{row})/L{row}, "")
    # Cột P: Kết Luận = =IF(O{row}="","",IF(O{row}<-0.02,"🔴 Bán phá giá sàn",IF(O{row}>0.05,"🔵 Bán cao hơn","🟢 Đúng giá sàn TMĐT")))
    
    f_diff_val = f'=IF(AND(ISNUMBER(M{curr_r}),ISNUMBER(L{curr_r})),M{curr_r}-L{curr_r},"")'
    f_diff_pct = f'=IF(AND(ISNUMBER(M{curr_r}),ISNUMBER(L{curr_r}),L{curr_r}>0),(M{curr_r}-L{curr_r})/L{curr_r},"")'
    f_status = f'=IF(O{curr_r}="","",IF(O{curr_r}<-0.02,"🔴 Bán phá giá sàn",IF(O{curr_r}>0.05,"🔵 Bán cao hơn","🟢 Đúng giá sàn TMĐT")))'
    
    row_vals = [
        idx,
        r["cust_code"],
        r["cust_name"],
        r["prov"],
        r["emp"],
        r["plat"],
        r["p_code"],
        r["p_name"],
        r["dvt"],
        r["gia_si"],
        r["gia_le"],
        r["gia_san_tmdt"],
        None, # Giá bán thực tế trên sàn: để trống để người dùng/nhân sự điền
        f_diff_val,
        f_diff_pct,
        f_status,
        None, # Ngày kiểm tra
        None, # Link bằng chứng
        None  # Ghi chú
    ]
    ws1.append(row_vals)
    
    for c_num in range(1, len(headers1) + 1):
        cell = ws1.cell(row=curr_r, column=c_num)
        cell.font = font_regular
        cell.border = thin_border
        if c_num in [1, 2, 4, 5, 6, 7, 9, 16, 17]:
            cell.alignment = align_center
        elif c_num in [10, 11, 12, 13, 14]:
            cell.alignment = align_right
            cell.number_format = "#,##0"
            if c_num == 12:
                cell.font = font_bold
        elif c_num == 15:
            cell.alignment = align_right
            cell.number_format = "0.0%"
            cell.font = font_bold
        else:
            cell.alignment = align_left

print("Tab 1 hoàn tất: 1,055 dòng có đầy đủ GIÁ SỈ, GIÁ LẺ và GIÁ SÀN TMĐT từ file PDF!")

# ==============================================================================
# TAB 2: TOÀN BỘ 1,289 SẢN PHẨM BẢNG BÁO GIÁ GỐC PDF (10.10.2026)
# ==============================================================================
ws2 = wb.create_sheet(title="2_Bang_Bao_Gia_Goc_PDF")
ws2.views.sheetView[0].showGridLines = True

headers2 = [
    "STT", "Trang PDF", "STT Gốc", "Tên / Mô Tả Sản Phẩm Trong Báo Giá", "ĐVT",
    "GIÁ SỈ (VNĐ)", "GIÁ LẺ (VNĐ)", "GIÁ SÀN TMĐT (VNĐ)", "Quy Cách / Ghi Chú"
]
ws2.append(headers2)
for col_num in range(1, len(headers2) + 1):
    c = ws2.cell(row=1, column=col_num)
    c.font = font_header
    c.fill = fill_accent
    if col_num == 8:
        c.fill = fill_gold
    c.alignment = align_center
    c.border = thin_border

for idx, p in enumerate(pdf_catalog, 1):
    curr_r = idx + 1
    row_vals = [
        idx,
        p["page"],
        p["stt"],
        p["item_desc"],
        p["dvt"],
        p["gia_si"],
        p["gia_le"],
        p["gia_san_tmdt"],
        p["suffix"]
    ]
    ws2.append(row_vals)
    for c_num in range(1, len(headers2) + 1):
        cell = ws2.cell(row=curr_r, column=c_num)
        cell.font = font_regular
        cell.border = thin_border
        if c_num in [1, 2, 3, 5]:
            cell.alignment = align_center
        elif c_num in [6, 7, 8]:
            cell.alignment = align_right
            cell.number_format = "#,##0"
            if c_num == 8:
                cell.font = font_bold
        else:
            cell.alignment = align_left

print("Tab 2 hoàn tất: 1,289 sản phẩm từ Bảng báo giá PDF 10.10.2026.")

# ==============================================================================
# TAB 3: DANH SÁCH 5 ĐẠI LÝ & SỐ LƯỢNG MÃ NHẬP
# ==============================================================================
ws3 = wb.create_sheet(title="3_Danh_Sach_Dai_Ly")
ws3.views.sheetView[0].showGridLines = True

headers3 = [
    "STT", "Mã Đại Lý", "Tên Khách Hàng / NPP / Đại Lý", "Tỉnh Thành",
    "Nhân Viên Phụ Trách", "Kênh Bán", "Số Mã Hàng Đã Nhập", "Trạng Thái Khảo Sát Giá Sàn"
]
ws3.append(headers3)
for col_num in range(1, len(headers3) + 1):
    c = ws3.cell(row=1, column=col_num)
    c.font = font_header
    c.fill = fill_header_gray
    c.alignment = align_center
    c.border = thin_border

dealer_summary = [
    (1, "KH001", "Điện nước Nhật (Shop) HCM", "HCM", "BẢO TRÂM", "Shopee", 257),
    (2, "KH002", "Nguyễn Hoàng Phú", "HCM", "BẢO TRÂM", "Shopee", 261),
    (3, "KH003", "CÔNG TY TNHH ĐỒ NGHỀ CẦM TAY PHƯƠNG MY", "HCM", "BẢO TRÂM", "Shopee", 165),
    (4, "KH004", "HCM-BIGSHOP", "HCM", "XUÂN TRANG", "Website bán hàng", 87),
    (5, "KH005", "ThinkSafe HCM", "HCM", "XUÂN TRANG", "Shopee", 285)
]

for d in dealer_summary:
    ws3.append([d[0], d[1], d[2], d[3], d[4], d[5], d[6], "Sẵn sàng nhập giá sàn thực tế để đối soát"])
    curr_r = d[0] + 1
    for c_num in range(1, len(headers3) + 1):
        cell = ws3.cell(row=curr_r, column=c_num)
        cell.font = font_regular
        cell.border = thin_border
        if c_num in [1, 2, 4, 5, 6, 8]:
            cell.alignment = align_center
        elif c_num == 7:
            cell.alignment = align_right
            cell.font = font_bold
        else:
            cell.alignment = align_left

# Autofit Column Widths
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if val.startswith("="):
                val = "12345678"
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

file_out = "KingBlue_Giam_Sat_Gia_San_TMDT_Chuan_PDF.xlsx"
wb.save(file_out)
print(f"🎉 Đã lưu file Excel hoàn chỉnh theo đúng PDF Báo giá tại: {file_out}")

