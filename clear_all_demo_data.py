import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE_PATH = os.path.join(BASE_DIR, "Bao_Cao_Cong_Viec_KingBlue.xlsx")
TOKEN_FILE = os.path.join(BASE_DIR, "API", "token.json")
SPREADSHEET_ID = "1_kID0uhutS6Ky_zpB2yW_AQXN2aUKKCKo_6tqZL1kbo"

# 9 nhân sự phòng Marketing King Blue
PERSONNEL_LIST = [
    {"id": "1", "tab_name": "Hoai Thuong - Content", "name": "Võ Thị Hoài Thương", "role": "Content Marketing & SEO Fanpage/Website", "group": "Content", "manager": "Marketing Manager"},
    {"id": "2", "tab_name": "Kieu Thuong - Content", "name": "Kiều Thương", "role": "Content Video & Hợp Tác KOC/Reviewer", "group": "Content", "manager": "Marketing Manager"},
    {"id": "3", "tab_name": "Thuy Thuong - Content", "name": "Thụy Thương", "role": "Trade Marketing & Chính Sách Điểm Bán", "group": "Content", "manager": "Marketing Manager"},
    {"id": "4", "tab_name": "Thien - Design", "name": "Thiện", "role": "Graphic Designer (2D/3D, POSM & Banner)", "group": "Design", "manager": "Marketing Manager"},
    {"id": "5", "tab_name": "Tu - Media", "name": "Tứ", "role": "Media / Photographer (Hình Ảnh, Video & User CRM)", "group": "Media", "manager": "Marketing Manager"},
    {"id": "6", "tab_name": "Thuc - San TMDT", "name": "Thức", "role": "Vận Hành Sàn TMĐT (Shopee & Lazada)", "group": "Sàn TMĐT", "manager": "Marketing Manager"},
    {"id": "7", "tab_name": "Ngan - San TMDT", "name": "Ngân", "role": "CSKH & Quản Trị Gian Hàng TikTok Shop", "group": "Sàn TMĐT", "manager": "Marketing Manager"},
    {"id": "8", "tab_name": "Hung - Kho MN", "name": "Hùng", "role": "Đóng Gói & Kho Vận Hàng Hóa Miền Nam", "group": "Đóng gói kho", "manager": "Marketing Manager"},
    {"id": "9", "tab_name": "Hung - Kho MB", "name": "Hùng Miền Bắc", "role": "Đóng Gói & Kho Vận Chi Nhánh Hà Nội", "group": "Đóng gói kho", "manager": "Marketing Manager"}
]

def clear_google_sheets():
    print("🧹 Đang xóa toàn bộ dữ liệu demo trên Google Sheets trực tuyến...")
    if not os.path.exists(TOKEN_FILE):
        print("⚠️ Không tìm thấy token.json!")
        return

    creds = Credentials.from_authorized_user_file(TOKEN_FILE)
    service = build("sheets", "v4", credentials=creds)

    # 1. Xóa sạch dữ liệu mẫu trong 9 tab cá nhân (Row 6 -> 50)
    clear_ranges = []
    blank_updates = []
    for p in PERSONNEL_LIST:
        t = p["tab_name"]
        clear_ranges.append(f"'{t}'!A6:K50")
        
        # Điền số thứ tự 1-15 cho các dòng trống sẵn sàng điền thật
        blank_rows = []
        for stt in range(1, 16):
            blank_rows.append([stt, "", "", "", "", "", "", "", "", ""])
        blank_updates.append({
            "range": f"'{t}'!A6:J20",
            "values": blank_rows
        })

    # Xóa dữ liệu mẫu tab BAO CAO TUAN
    clear_ranges.append("'BAO CAO TUAN'!A5:H20")

    # Thực hiện clear
    service.spreadsheets().values().batchClear(
        spreadsheetId=SPREADSHEET_ID,
        body={"ranges": clear_ranges}
    ).execute()
    print("✅ Đã xóa sạch dữ liệu mẫu của 9 sheet nhân sự và sheet báo cáo tuần.")

    # Điền khung mẫu trống
    service.spreadsheets().values().batchUpdate(
        spreadsheetId=SPREADSHEET_ID,
        body={
            "valueInputOption": "USER_ENTERED",
            "data": blank_updates
        }
    ).execute()
    print("✅ Đã thiết lập khung dòng trống chuẩn bị nộp thật cho 9 nhân sự.")

    # 2. Cài đặt lại công thức động trên sheet TONG HOP HANG NGAY
    print("🔄 Cập nhật lại công thức tự động cho sheet TONG HOP HANG NGAY...")
    date_cells = [
        ["📅 NGÀY XEM BÁO CÁO:", "05/10/2026", "(Gõ ngày cần xem theo định dạng DD/MM/YYYY hoặc =TODAY())"]
    ]
    service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID,
        range="'TONG HOP HANG NGAY'!B3:D3",
        valueInputOption="USER_ENTERED",
        body={"values": date_cells}
    ).execute()

    kpi_formulas = [
        ['=COUNTIF(J8:J16; "Đúng hạn") & " / 9 (" & TEXT(COUNTIF(J8:J16; "Đúng hạn")/9; "0%") & ")"',
         '',
         '=COUNTIF(F8:F16; "*•*") & " Vấn đề"',
         '',
         '=IF(COUNTIF(J8:J16; "Chưa nộp")=0; "100% Hoàn thành"; "Còn " & COUNTIF(J8:J16; "Chưa nộp") & " chưa nộp")']
    ]
    service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID,
        range="'TONG HOP HANG NGAY'!D5:H5",
        valueInputOption="USER_ENTERED",
        body={"values": kpi_formulas}
    ).execute()

    master_formulas = []
    for idx, p in enumerate(PERSONNEL_LIST, start=8):
        t = p["tab_name"]
        f_res = f'=IFERROR(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$D$6:$D$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$D$6:$D$50; "⏳ Chưa nộp"); "⏳ Chưa nộp")); "⏳ Chưa nộp")'
        f_diff = f'=IFERROR(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$E$6:$E$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$E$6:$E$50; "-"); "-")); "-")'
        f_lesson = f'=IFERROR(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$F$6:$F$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$F$6:$F$50; "-"); "-")); "-")'
        f_plan = f'=IFERROR(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$G$6:$G$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$G$6:$G$50; "-"); "-")); "-")'
        f_time = f'=IFERROR(TEXT(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$C$6:$C$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$C$6:$C$50; "--:--"); "--:--")); "hh:mm:ss"); "--:--")'
        f_status = f'=IF(E{idx}="⏳ Chưa nộp"; "Chưa nộp"; IFERROR(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$I$6:$I$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$I$6:$I$50; "Đúng hạn"); "Đúng hạn")); "Đúng hạn"))'
        f_feedback = f'=IFERROR(XLOOKUP(TEXT($C$3; "dd/mm/yyyy"); \'{t}\'!$B$6:$B$50; \'{t}\'!$J$6:$J$50; IFERROR(XLOOKUP($C$3; \'{t}\'!$B$6:$B$50; \'{t}\'!$J$6:$J$50; "-"); "-")); "-")'

        master_formulas.append({
            "range": f"'TONG HOP HANG NGAY'!E{idx}:K{idx}",
            "values": [[f_res, f_diff, f_lesson, f_plan, f_time, f_status, f_feedback]]
        })

    service.spreadsheets().values().batchUpdate(
        spreadsheetId=SPREADSHEET_ID,
        body={
            "valueInputOption": "USER_ENTERED",
            "data": master_formulas
        }
    ).execute()
    print("✅ Đã kết nối công thức sạch cho sheet TONG HOP HANG NGAY.")

def rebuild_clean_excel_file():
    print("📁 Đang tạo lại file Excel Bao_Cao_Cong_Viec_KingBlue.xlsx hoàn toàn sạch dữ liệu demo...")
    wb = openpyxl.Workbook()

    c_primary = "1A365D"
    c_accent = "2B6CB0"
    c_light_blue = "EBF8FF"
    c_header_gray = "EDF2F7"
    c_white = "FFFFFF"
    c_border = "CBD5E0"

    font_title = Font(name="Calibri", size=15, bold=True, color=c_white)
    font_header = Font(name="Calibri", size=10, bold=True, color=c_white)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_kpi_num = Font(name="Calibri", size=16, bold=True, color=c_primary)
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="4A5568")

    fill_primary = PatternFill(start_color=c_primary, end_color=c_primary, fill_type="solid")
    fill_accent = PatternFill(start_color=c_accent, end_color=c_accent, fill_type="solid")
    fill_kpi = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")
    fill_sub = PatternFill(start_color=c_header_gray, end_color=c_header_gray, fill_type="solid")
    fill_pending = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=c_border),
        right=Side(style='thin', color=c_border),
        top=Side(style='thin', color=c_border),
        bottom=Side(style='thin', color=c_border)
    )
    kpi_border = Border(
        left=Side(style='medium', color=c_accent),
        right=Side(style='medium', color=c_accent),
        top=Side(style='medium', color=c_accent),
        bottom=Side(style='medium', color=c_accent)
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_top_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # 1. TAB MASTER: TONG HOP HANG NGAY
    ws_master = wb.active
    ws_master.title = "TONG HOP HANG NGAY"
    ws_master.views.sheetView[0].showGridLines = True

    ws_master.merge_cells("A1:K2")
    ws_master["A1"] = "BẢNG TỔNG HỢP BÁO CÁO CÔNG VIỆC HÀNG NGÀY - PHÒNG MARKETING KING BLUE"
    ws_master["A1"].font = font_title
    ws_master["A1"].fill = fill_primary
    ws_master["A1"].alignment = align_center

    ws_master["B3"] = "📅 NGÀY XEM BÁO CÁO:"
    ws_master["B3"].font = font_bold
    ws_master["C3"] = "05/10/2026"
    ws_master["C3"].font = font_bold
    ws_master["D3"] = "(Gõ ngày cần xem hoặc gõ =TODAY())"

    kpi_cards = [
        ("B4:C4", "B5:C5", "TỔNG NHÂN SỰ", "9 Người"),
        ("D4:E4", "D5:E5", "ĐÃ NỘP HÔM NAY", '0 / 9 (0%)'),
        ("F4:G4", "F5:G5", "VƯỚNG MẮC PHÁT SINH", "0 Vấn đề"),
        ("H4:I4", "H5:I5", "TÌNH TRẠNG NỘP", "Chờ nộp báo cáo"),
        ("J4:K4", "J5:K5", "TRƯỞNG PHÒNG", "Marketing Manager"),
    ]
    for top_range, bot_range, label, val in kpi_cards:
        ws_master.merge_cells(top_range)
        ws_master.merge_cells(bot_range)
        top_cell = ws_master[top_range.split(":")[0]]
        bot_cell = ws_master[bot_range.split(":")[0]]
        top_cell.value = label
        top_cell.font = font_kpi_label
        top_cell.fill = fill_sub
        top_cell.alignment = align_center
        bot_cell.value = val
        bot_cell.font = font_kpi_num
        bot_cell.fill = fill_kpi
        bot_cell.alignment = align_center
        for row in ws_master[f"{top_range.split(':')[0]}:{bot_range.split(':')[1]}"]:
            for cell in row:
                cell.border = kpi_border

    headers_master = [
        "STT", "Nhân Sự", "Bộ Phận", "Vị Trí Chuyên Môn",
        "1️⃣ KẾT QUẢ ĐẠT ĐƯỢC HÔM NAY",
        "2️⃣ KHÓ KHĂN / VƯỚNG MẮC",
        "3️⃣ BÀI HỌC / ĐỀ XUẤT",
        "4️⃣ KẾ HOẠCH NGÀY MAI",
        "Giờ Nộp", "Tình Trạng", "Marketing Manager Phản Hồi"
    ]
    ws_master.row_dimensions[7].height = 28
    for col_idx, h_text in enumerate(headers_master, start=1):
        c = ws_master.cell(row=7, column=col_idx, value=h_text)
        c.font = font_header
        c.fill = fill_accent
        c.alignment = align_center
        c.border = thin_border

    # Điền 9 dòng nhân sự ở trạng thái chờ nộp sạch sẽ
    for idx, p in enumerate(PERSONNEL_LIST, start=8):
        ws_master.row_dimensions[idx].height = 36
        vals = [
            int(p["id"]), p["name"], p["group"], p["role"],
            "⏳ Chưa nộp", "-", "-", "-", "--:--", "Chưa nộp", "-"
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws_master.cell(row=idx, column=c_idx, value=v)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx in [1, 9, 10]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            if c_idx == 10:
                cell.fill = fill_pending
                cell.font = font_bold

    col_widths_master = {
        'A': 6, 'B': 20, 'C': 14, 'D': 25, 'E': 42,
        'F': 32, 'G': 32, 'H': 35, 'I': 10, 'J': 13, 'K': 30
    }
    for col_letter, width in col_widths_master.items():
        ws_master.column_dimensions[col_letter].width = width

    # 2. TẠO 9 TAB NHÂN SỰ HOÀN TOÀN TRẮNG
    headers_ind = [
        "STT", "Ngày Báo Cáo", "Giờ Nộp",
        "1️⃣ KẾT QUẢ ĐẠT ĐƯỢC (Nhiệm vụ 1, 2, 3...)",
        "2️⃣ KHÓ KHĂN / VƯỚNG MẮC",
        "3️⃣ BÀI HỌC KINH NGHIỆM / ĐỀ XUẤT",
        "4️⃣ KẾ HOẠCH CÔNG VIỆC NGÀY MAI",
        "Link Minh Chứng", "Tình Trạng", "Marketing Manager Phê Duyệt & Phản Hồi"
    ]

    for p in PERSONNEL_LIST:
        ws_p = wb.create_sheet(title=p["tab_name"])
        ws_p.views.sheetView[0].showGridLines = True

        ws_p.merge_cells("A1:J2")
        ws_p["A1"] = f"NHẬT KÝ BÁO CÁO CÔNG VIỆC HÀNG NGÀY - {p['name'].upper()}"
        ws_p["A1"].font = font_title
        ws_p["A1"].fill = fill_primary
        ws_p["A1"].alignment = align_center

        ws_p.merge_cells("A3:B3")
        ws_p["A3"] = "Họ và Tên:"
        ws_p["A3"].font = font_bold
        ws_p.merge_cells("C3:D3")
        ws_p["C3"] = p["name"]
        ws_p["C3"].font = font_bold

        ws_p.merge_cells("E3:F3")
        ws_p["E3"] = "Vị trí chuyên môn:"
        ws_p["E3"].font = font_bold
        ws_p.merge_cells("G3:H3")
        ws_p["G3"] = p["role"]

        ws_p["I3"] = "Quản lý duyệt:"
        ws_p["I3"].font = font_bold
        ws_p["J3"] = p["manager"]
        ws_p["J3"].font = font_bold

        for row in ws_p["A3:J3"]:
            for c in row:
                c.fill = fill_sub
                c.border = thin_border

        ws_p.row_dimensions[5].height = 26
        for col_idx, h_text in enumerate(headers_ind, start=1):
            c = ws_p.cell(row=5, column=col_idx, value=h_text)
            c.font = font_header
            c.fill = fill_accent
            c.alignment = align_center
            c.border = thin_border

        # 15 dòng trống định dạng sẵn
        for r_idx in range(6, 21):
            ws_p.row_dimensions[r_idx].height = 32
            ws_p.cell(row=r_idx, column=1, value=r_idx - 5).alignment = align_top_center
            for c_idx in range(1, 11):
                cell = ws_p.cell(row=r_idx, column=c_idx)
                cell.border = thin_border

        col_widths_ind = {
            'A': 6, 'B': 13, 'C': 10, 'D': 44, 'E': 34,
            'F': 34, 'G': 36, 'H': 20, 'I': 13, 'J': 32
        }
        for col_letter, width in col_widths_ind.items():
            ws_p.column_dimensions[col_letter].width = width

    # 3. TAB BAO CAO TUAN
    ws_w = wb.create_sheet(title="BAO CAO TUAN")
    ws_w.views.sheetView[0].showGridLines = True
    ws_w.merge_cells("A1:H2")
    ws_w["A1"] = "BÁO CÁO TỔNG KẾT TUẦN - PHÒNG MARKETING KING BLUE"
    ws_w["A1"].font = font_title
    ws_w["A1"].fill = fill_primary
    ws_w["A1"].alignment = align_center

    headers_w = [
        "STT", "Họ và Tên", "Vị Trí", "Mục Tiêu Tuần",
        "Kết Quả Đạt Được (% Hoàn thành)", "Vấn Đề Tồn Đọng", "Kế Hoạch Tuần Tới", "Marketing Manager Đánh Giá"
    ]
    ws_w.row_dimensions[4].height = 25
    for col_idx, h_text in enumerate(headers_w, start=1):
        c = ws_w.cell(row=4, column=col_idx, value=h_text)
        c.font = font_header
        c.fill = fill_accent
        c.alignment = align_center
        c.border = thin_border

    for idx, p in enumerate(PERSONNEL_LIST, start=5):
        ws_w.row_dimensions[idx].height = 32
        ws_w.cell(row=idx, column=1, value=int(p["id"])).alignment = align_center
        ws_w.cell(row=idx, column=2, value=p["name"])
        ws_w.cell(row=idx, column=3, value=p["role"])
        for c_idx in range(1, 9):
            ws_w.cell(row=idx, column=c_idx).border = thin_border

    col_widths_w = {'A': 6, 'B': 22, 'C': 22, 'D': 35, 'E': 20, 'F': 25, 'G': 32, 'H': 25}
    for col_letter, width in col_widths_w.items():
        ws_w.column_dimensions[col_letter].width = width

    wb.save(FILE_PATH)
    print(f"✅ Đã lưu file Excel sạch dữ liệu demo tại: {FILE_PATH}")

if __name__ == "__main__":
    clear_google_sheets()
    rebuild_clean_excel_file()
