import openpyxl
import glob
import json
import os
from datetime import datetime, timedelta
from collections import defaultdict, Counter

def main():
    files = glob.glob('/Users/Admin/Documents/KIng BLue/**/So_chi_tiet_ban_hang 2909.xlsx', recursive=True)
    if not files:
        print("Error: Excel file not found!")
        return
    filepath = files[0]
    print(f"Loading data from: {filepath}")

    wb = openpyxl.load_workbook(filepath, read_only=True)

    customers = {}
    products = defaultdict(lambda: {
        'code': '', 'name': '', 'unit': 'Cái', 'category': 'Khác',
        'qty': 0.0, 'promo_qty': 0.0, 'gross': 0.0, 'discount': 0.0,
        'return_qty': 0.0, 'return_val': 0.0, 'net': 0.0,
        'orders': set(), 'customers': set(), 'prices': [],
        'monthly': defaultdict(float), 'monthly_qty': defaultdict(float),
        'top_customers': defaultdict(float)
    })
    
    monthly_stats = defaultdict(lambda: {'gross': 0, 'discount': 0, 'return': 0, 'net': 0, 'orders': set(), 'customers': set(), 'items': 0})
    province_stats = defaultdict(lambda: {'gross': 0, 'discount': 0, 'return': 0, 'net': 0, 'orders': set(), 'customers': set()})
    sales_rep_stats = defaultdict(lambda: {'gross': 0, 'discount': 0, 'return': 0, 'net': 0, 'orders': set(), 'customers': set()})
    category_stats = defaultdict(lambda: {'gross': 0, 'qty': 0, 'skus': set(), 'customers': set(), 'orders': set()})
    cust_order_dates = defaultdict(set)

    ref_date = datetime(2026, 9, 30)

    total_gross = 0.0
    total_discount = 0.0
    total_return = 0.0
    total_net = 0.0
    all_orders = set()
    total_line_items = 0
    total_qty = 0.0

    sheet_names = ['SỔ CHI TIẾT BÁN HÀNG', 'SỔ CHI TIẾT BÁN HÀNG_1']

    for sname in sheet_names:
        if sname not in wb.sheetnames: continue
        ws = wb[sname]
        print(f"Reading sheet: {sname}...")
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r or r[0] is None or r[3] is None: continue
            
            total_line_items += 1
            raw_date = r[0]
            if isinstance(raw_date, datetime): dt = raw_date
            else:
                try: dt = datetime.strptime(str(raw_date)[:10], '%Y-%m-%d')
                except: dt = datetime(2026, 1, 1)
            
            month_key = dt.strftime('%Y-%m')
            so_ct = str(r[1]).strip() if r[1] is not None else ''
            ma_kh = str(r[3]).strip()
            ten_kh = str(r[4]).strip() if r[4] is not None else ma_kh
            dia_chi = str(r[5]).strip() if r[5] is not None else ''
            sdt = str(r[6]).strip() if r[6] is not None else ''
            ma_hang = str(r[7]).strip() if r[7] is not None else ''
            ten_hang = str(r[8]).strip() if r[8] is not None else ma_hang
            sl_ban = float(r[9] or 0)
            dvt = str(r[11]).strip() if r[11] is not None else 'Cái'
            sl_km = float(r[12] or 0)
            don_gia = float(r[15] or 0)
            doanh_so = float(r[19] or 0)
            chiet_khau = float(r[24] or 0)
            sl_tra = float(r[25] or 0)
            tra_lai = float(r[32] or 0)
            nvbh = str(r[38]).strip() if r[38] is not None else 'Chưa gán'
            tinh = str(r[39]).strip() if r[39] is not None else 'Chưa xác định'
            nhom = str(r[42]).strip() if r[42] is not None else 'Khác'
            
            if not tinh or tinh == 'None': tinh = 'Chưa xác định'
            if not nvbh or nvbh == 'None': nvbh = 'Chưa gán'
            if not nhom or nhom == 'None': nhom = 'Khác'
            if not dvt or dvt == 'None': dvt = 'Cái'

            net_rev = doanh_so - chiet_khau - tra_lai
            
            total_gross += doanh_so
            total_discount += chiet_khau
            total_return += tra_lai
            total_net += net_rev
            total_qty += sl_ban
            if so_ct: all_orders.add(so_ct)
            if dt: cust_order_dates[ma_kh].add(dt.date())
                
            m = monthly_stats[month_key]
            m['gross'] += doanh_so
            m['discount'] += chiet_khau
            m['return'] += tra_lai
            m['net'] += net_rev
            m['items'] += sl_ban
            if so_ct: m['orders'].add(so_ct)
            m['customers'].add(ma_kh)
            
            p = province_stats[tinh]
            p['gross'] += doanh_so
            p['discount'] += chiet_khau
            p['return'] += tra_lai
            p['net'] += net_rev
            if so_ct: p['orders'].add(so_ct)
            p['customers'].add(ma_kh)
            
            s = sales_rep_stats[nvbh]
            s['gross'] += doanh_so
            s['discount'] += chiet_khau
            s['return'] += tra_lai
            s['net'] += net_rev
            if so_ct: s['orders'].add(so_ct)
            s['customers'].add(ma_kh)
            
            c = category_stats[nhom]
            c['gross'] += doanh_so
            c['qty'] += sl_ban
            c['skus'].add(ma_hang)
            c['customers'].add(ma_kh)
            if so_ct: c['orders'].add(so_ct)
            
            # Product level
            pr = products[ma_hang]
            pr['code'] = ma_hang
            pr['name'] = ten_hang
            pr['unit'] = dvt
            pr['category'] = nhom
            pr['qty'] += sl_ban
            pr['promo_qty'] += sl_km
            pr['gross'] += doanh_so
            pr['discount'] += chiet_khau
            pr['return_qty'] += sl_tra
            pr['return_val'] += tra_lai
            pr['net'] += net_rev
            if so_ct: pr['orders'].add(so_ct)
            if ma_kh: pr['customers'].add(ma_kh)
            if don_gia > 0: pr['prices'].append(don_gia)
            pr['monthly'][month_key] += doanh_so
            pr['monthly_qty'][month_key] += sl_ban
            pr['top_customers'][ten_kh] += doanh_so
            
            # Customer level
            if ma_kh not in customers:
                customers[ma_kh] = {
                    'id': ma_kh,
                    'name': ten_kh,
                    'address': dia_chi,
                    'phone': sdt,
                    'province': tinh,
                    'sales_reps': defaultdict(int),
                    'first_date': dt,
                    'last_date': dt,
                    'orders': set(),
                    'line_items': 0,
                    'total_qty': 0,
                    'gross': 0.0,
                    'discount': 0.0,
                    'return': 0.0,
                    'net': 0.0,
                    'categories': defaultdict(float),
                    'top_products': defaultdict(lambda: {'name': '', 'gross': 0, 'qty': 0})
                }
            
            cust = customers[ma_kh]
            if dt < cust['first_date']: cust['first_date'] = dt
            if dt > cust['last_date']: cust['last_date'] = dt
            if so_ct: cust['orders'].add(so_ct)
            cust['line_items'] += 1
            cust['total_qty'] += sl_ban
            cust['gross'] += doanh_so
            cust['discount'] += chiet_khau
            cust['return'] += tra_lai
            cust['net'] += net_rev
            cust['sales_reps'][nvbh] += 1
            cust['categories'][nhom] += doanh_so
            
            prod_entry = cust['top_products'][ma_hang]
            prod_entry['name'] = ten_hang
            prod_entry['gross'] += doanh_so
            prod_entry['qty'] += sl_ban

    # 1. Process Customers
    sorted_customers = sorted(customers.values(), key=lambda x: x['net'], reverse=True)
    cum_rev = 0.0
    for c in sorted_customers:
        cum_rev += c['net']
        c_pct = cum_rev / total_net if total_net > 0 else 0
        if c_pct <= 0.80: c['abc'] = 'A'
        elif c_pct <= 0.95: c['abc'] = 'B'
        else: c['abc'] = 'C'

        r_days = (ref_date - c['last_date']).days
        f_orders = len(c['orders'])
        m_val = c['net']
        c['r_days'] = r_days
        c['f_orders'] = f_orders
        c['main_sales_rep'] = max(c['sales_reps'].items(), key=lambda x: x[1])[0] if c['sales_reps'] else 'Chưa gán'

        if r_days <= 30 and (f_orders >= 15 or m_val >= 300_000_000):
            c['segment'] = 'VIP / Khách hàng Vàng'
            c['segment_code'] = 'CHAMPION'
            c['segment_color'] = '#3b82f6'
            c['action'] = 'Chăm sóc đặc biệt, ký hợp đồng năm, tặng quà tri ân VIP, ưu tiên giữ hàng.'
        elif r_days <= 45 and (f_orders >= 8 or m_val >= 100_000_000):
            c['segment'] = 'Khách hàng Trung thành'
            c['segment_code'] = 'LOYAL'
            c['segment_color'] = '#10b981'
            c['action'] = 'Upsell các dòng máy mới, duy trì tần suất gọi điện định kỳ 1 tuần/lần.'
        elif r_days <= 45 and (f_orders >= 2 or m_val >= 30_000_000):
            c['segment'] = 'Tiềm năng phát triển'
            c['segment_code'] = 'POTENTIAL'
            c['segment_color'] = '#06b6d4'
            c['action'] = 'Đề xuất chính sách chiết khấu bậc thang theo doanh số tháng để kích cầu.'
        elif r_days <= 45:
            c['segment'] = 'Khách hàng Mới / Mua gần đây'
            c['segment_code'] = 'NEW_ACTIVE'
            c['segment_color'] = '#8b5cf6'
            c['action'] = 'Theo dõi trải nghiệm sử dụng đợt đầu, gửi catalogue giới thiệu thêm mã hàng.'
        elif 45 < r_days <= 90 and (f_orders >= 5 or m_val >= 50_000_000):
            c['segment'] = 'Cần chăm sóc / Nguy cơ nguội'
            c['segment_code'] = 'NEEDS_ATTENTION'
            c['segment_color'] = '#f59e0b'
            c['action'] = 'Sales phụ trách cần liên hệ hỏi thăm tồn kho, giới thiệu mã hàng khuyến mãi.'
        elif 90 < r_days <= 150:
            c['segment'] = 'Nguy cơ rời bỏ cao'
            c['segment_code'] = 'AT_RISK'
            c['segment_color'] = '#ef4444'
            c['action'] = 'BÁO ĐỘNG: Gặp trực tiếp chủ đại lý, tìm hiểu lý do ngưng nhập hàng, chào chính sách đặc biệt.'
        else:
            c['segment'] = 'Ngủ đông / Rời bỏ'
            c['segment_code'] = 'LOST'
            c['segment_color'] = '#64748b'
            c['action'] = 'Chiến dịch re-engagement: gửi thông báo sản phẩm mới kèm voucher ưu đãi mở lại tài khoản.'

    # Reorder Prediction Analysis
    reorder_predictions = []
    for c in sorted_customers:
        ma_kh = c['id']
        dates = sorted(list(cust_order_dates[ma_kh]))
        top_prods = sorted(c['top_products'].values(), key=lambda x: x['gross'], reverse=True)[:3]
        top_prod_names = [p['name'] for p in top_prods]
        
        last_d = c['last_date'].date()
        r_days = c['r_days']
        
        if len(dates) >= 2:
            intervals = [(dates[i] - dates[i-1]).days for i in range(1, len(dates))]
            avg_cycle = sum(intervals) / len(intervals)
            days_left = round(avg_cycle - r_days, 1)
            pred_date = last_d + timedelta(days=round(avg_cycle))
            depletion_pct = round((r_days / avg_cycle) * 100, 1) if avg_cycle > 0 else 100.0
            
            if r_days > avg_cycle * 1.3:
                status = 'OVERDUE'
                status_label = 'Đã Quá Hạn Chu Kỳ'
                status_color = '#ef4444' # đỏ
                urgency = 'Báo động: Đã quá chu kỳ bình thường. Nguy cơ đại lý cạn hàng hoặc chuyển sang mua đối thủ!'
                score = 95
            elif r_days >= avg_cycle * 0.8:
                status = 'RUNNING_LOW'
                status_label = 'Sắp Hết Hàng (Cần Gọi Ngay)'
                status_color = '#f97316' # cam
                urgency = 'Khẩn cấp: Tồn kho đại lý dự kiến cạn trong 1-5 ngày tới. Thời điểm vàng để chốt đơn gối đầu!'
                score = 90
            elif r_days >= avg_cycle * 0.5:
                status = 'UPCOMING'
                status_label = 'Sắp Đến Chu Kỳ'
                status_color = '#eab308' # vàng
                urgency = 'Chuẩn bị: Dự kiến cần nhập hàng trong 7-15 ngày tới. Gửi chương trình khuyến mãi tháng mới.'
                score = 50
            else:
                status = 'SAFE'
                status_label = 'Tồn Kho An Toàn'
                status_color = '#10b981' # xanh lá
                urgency = 'An toàn: Mới nhập đợt hàng gần đây, tồn kho tại cửa hàng còn dồi dào.'
                score = 20
        else:
            avg_cycle = 30.0
            days_left = round(30.0 - r_days, 1)
            pred_date = last_d + timedelta(days=30)
            depletion_pct = round((r_days / 30.0) * 100, 1)
            if r_days > 45:
                status = 'OVERDUE'
                status_label = 'Đã Lâu Chưa Mua Lại'
                status_color = '#ef4444'
                urgency = 'Khách mới mua 1 lần nhưng đã quá 45 ngày chưa quay lại.'
                score = 60
            else:
                status = 'UPCOMING'
                status_label = 'Khách Mới Đang Sử Dụng'
                status_color = '#eab308'
                urgency = 'Khách mới mua 1 lần, cần gọi điện khảo sát mức độ hài lòng.'
                score = 40

        reorder_predictions.append({
            'id': c['id'],
            'name': c['name'],
            'phone': c['phone'],
            'province': c['province'],
            'main_sales_rep': c['main_sales_rep'],
            'net': round(c['net'], 0),
            'gross': round(c['gross'], 0),
            'orders_count': len(c['orders']),
            'last_date': c['last_date'].strftime('%Y-%m-%d'),
            'r_days': c['r_days'],
            'avg_cycle': round(avg_cycle, 1),
            'days_left': days_left,
            'predicted_date': pred_date.strftime('%Y-%m-%d'),
            'depletion_pct': depletion_pct,
            'status': status,
            'status_label': status_label,
            'status_color': status_color,
            'urgency': urgency,
            'score': score,
            'abc': c['abc'],
            'top_products': top_prod_names
        })

    # Sort prediction: High score and high revenue first
    reorder_predictions_sorted = sorted(reorder_predictions, key=lambda x: (x['score'], x['net']), reverse=True)

    # Prediction KPI stats
    pred_status_counts = Counter([p['status'] for p in reorder_predictions])
    pred_status_rev = defaultdict(float)
    for p in reorder_predictions: pred_status_rev[p['status']] += p['net']

    prediction_summary = {
        'running_low_count': pred_status_counts['RUNNING_LOW'],
        'running_low_rev': round(pred_status_rev['RUNNING_LOW'], 0),
        'overdue_count': pred_status_counts['OVERDUE'],
        'overdue_rev': round(pred_status_rev['OVERDUE'], 0),
        'upcoming_count': pred_status_counts['UPCOMING'],
        'upcoming_rev': round(pred_status_rev['UPCOMING'], 0),
        'safe_count': pred_status_counts['SAFE'],
        'safe_rev': round(pred_status_rev['SAFE'], 0)
    }

    # Customers list
    formatted_customers = []
    for c in sorted_customers:
        top_cats = sorted(c['categories'].items(), key=lambda x: x[1], reverse=True)[:5]
        top_prods = sorted(c['top_products'].values(), key=lambda x: x['gross'], reverse=True)[:5]
        formatted_customers.append({
            'id': c['id'],
            'name': c['name'],
            'address': c['address'],
            'phone': c['phone'],
            'province': c['province'],
            'main_sales_rep': c['main_sales_rep'],
            'first_date': c['first_date'].strftime('%Y-%m-%d'),
            'last_date': c['last_date'].strftime('%Y-%m-%d'),
            'r_days': c['r_days'],
            'f_orders': c['f_orders'],
            'line_items': c['line_items'],
            'total_qty': round(c['total_qty'], 1),
            'gross': round(c['gross'], 0),
            'discount': round(c['discount'], 0),
            'discount_pct': round((c['discount'] / c['gross'] * 100) if c['gross'] > 0 else 0, 2),
            'return': round(c['return'], 0),
            'net': round(c['net'], 0),
            'abc': c['abc'],
            'segment': c['segment'],
            'segment_code': c['segment_code'],
            'segment_color': c['segment_color'],
            'action': c['action'],
            'top_categories': [{'name': cat, 'gross': round(val, 0)} for cat, val in top_cats],
            'top_products': [{'name': pr['name'], 'gross': round(pr['gross'], 0), 'qty': pr['qty']} for pr in top_prods]
        })

    # Products list
    sorted_products = sorted(products.values(), key=lambda x: x['gross'], reverse=True)
    cum_prod_rev = 0.0
    for pr in sorted_products:
        cum_prod_rev += pr['gross']
        pct = cum_prod_rev / total_gross if total_gross > 0 else 0
        if pct <= 0.80:
            pr['abc'] = 'A'
            pr['abc_label'] = 'Hạng A (Chủ lực 80%)'
            pr['abc_color'] = '#2563eb'
            pr['action'] = 'Duy trì tồn kho tối đa, ưu tiên sản xuất/nhập khẩu, không được để thiếu hàng.'
        elif pct <= 0.95:
            pr['abc'] = 'B'
            pr['abc_label'] = 'Hạng B (Tiềm năng 15%)'
            pr['abc_color'] = '#0ea5e9'
            pr['action'] = 'Đóng combo khuyến mãi cùng sản phẩm Hạng A để kích thích đại lý nhập thêm.'
        else:
            pr['abc'] = 'C'
            pr['abc_label'] = 'Hạng C (Chậm luân chuyển 5%)'
            pr['abc_color'] = '#94a3b8'
            pr['action'] = 'Rà soát danh mục: xem xét xả hàng tồn kho hoặc ngừng nhập nếu tỷ lệ quay vòng quá chậm.'

    formatted_products = []
    prod_abc_summary = {'A': {'count': 0, 'gross': 0, 'qty': 0}, 'B': {'count': 0, 'gross': 0, 'qty': 0}, 'C': {'count': 0, 'gross': 0, 'qty': 0}}
    
    for pr in sorted_products:
        avg_price = sum(pr['prices']) / len(pr['prices']) if pr['prices'] else 0
        min_price = min(pr['prices']) if pr['prices'] else 0
        max_price = max(pr['prices']) if pr['prices'] else 0
        
        prod_abc_summary[pr['abc']]['count'] += 1
        prod_abc_summary[pr['abc']]['gross'] += pr['gross']
        prod_abc_summary[pr['abc']]['qty'] += pr['qty']

        top_custs = sorted(pr['top_customers'].items(), key=lambda x: x[1], reverse=True)[:5]
        monthly_arr = [{'month': m, 'gross': round(pr['monthly'][m], 0), 'qty': round(pr['monthly_qty'][m], 0)} for m in sorted(monthly_stats.keys())]

        formatted_products.append({
            'code': pr['code'],
            'name': pr['name'],
            'unit': pr['unit'],
            'category': pr['category'],
            'qty': round(pr['qty'], 1),
            'promo_qty': round(pr['promo_qty'], 1),
            'gross': round(pr['gross'], 0),
            'discount': round(pr['discount'], 0),
            'discount_pct': round((pr['discount'] / pr['gross'] * 100) if pr['gross'] > 0 else 0, 2),
            'return_qty': round(pr['return_qty'], 1),
            'return_val': round(pr['return_val'], 0),
            'return_pct': round((pr['return_val'] / pr['gross'] * 100) if pr['gross'] > 0 else 0, 2),
            'net': round(pr['net'], 0),
            'avg_price': round(avg_price, 0),
            'min_price': round(min_price, 0),
            'max_price': round(max_price, 0),
            'orders_count': len(pr['orders']),
            'customers_count': len(pr['customers']),
            'customer_penetration_pct': round((len(pr['customers']) / len(customers) * 100), 1),
            'abc': pr['abc'],
            'abc_label': pr['abc_label'],
            'abc_color': pr['abc_color'],
            'action': pr['action'],
            'top_customers': [{'name': c_name, 'gross': round(c_gross, 0)} for c_name, c_gross in top_custs],
            'monthly': monthly_arr
        })

    # Summary KPIs
    kpi_summary = {
        'total_gross': round(total_gross, 0),
        'total_discount': round(total_discount, 0),
        'discount_pct': round((total_discount / total_gross * 100) if total_gross > 0 else 0, 2),
        'total_return': round(total_return, 0),
        'return_pct': round((total_return / total_gross * 100) if total_gross > 0 else 0, 2),
        'total_net': round(total_net, 0),
        'total_orders': len(all_orders),
        'total_customers': len(customers),
        'total_skus': len(products),
        'total_line_items': total_line_items,
        'total_qty': round(total_qty, 0),
        'aov': round(total_net / len(all_orders) if all_orders else 0, 0),
        'arpu': round(total_net / len(customers) if customers else 0, 0),
        'avg_orders_per_cust': round(len(all_orders) / len(customers) if customers else 0, 1),
        'date_range': '01/01/2026 - 29/09/2026'
    }

    monthly_data = []
    for m in sorted(monthly_stats.keys()):
        v = monthly_stats[m]
        monthly_data.append({
            'month': m,
            'month_label': f"Tháng {int(m.split('-')[1])}",
            'gross': round(v['gross'], 0),
            'discount': round(v['discount'], 0),
            'discount_pct': round((v['discount'] / v['gross'] * 100) if v['gross'] > 0 else 0, 1),
            'return': round(v['return'], 0),
            'net': round(v['net'], 0),
            'orders': len(v['orders']),
            'customers': len(v['customers']),
            'items': round(v['items'], 0)
        })

    province_data = []
    sorted_prov = sorted(province_stats.items(), key=lambda x: x[1]['net'], reverse=True)
    for p, v in sorted_prov:
        province_data.append({
            'province': p,
            'net': round(v['net'], 0),
            'gross': round(v['gross'], 0),
            'discount': round(v['discount'], 0),
            'return': round(v['return'], 0),
            'customers': len(v['customers']),
            'orders': len(v['orders']),
            'pct_of_total': round((v['net'] / total_net * 100) if total_net > 0 else 0, 2)
        })

    sales_rep_data = []
    sorted_sales = sorted(sales_rep_stats.items(), key=lambda x: x[1]['net'], reverse=True)
    for s, v in sorted_sales:
        sales_rep_data.append({
            'name': s,
            'net': round(v['net'], 0),
            'gross': round(v['gross'], 0),
            'discount': round(v['discount'], 0),
            'customers': len(v['customers']),
            'orders': len(v['orders']),
            'arpu': round((v['net'] / len(v['customers'])) if v['customers'] else 0, 0),
            'pct_of_total': round((v['net'] / total_net * 100) if total_net > 0 else 0, 2)
        })

    category_data = []
    sorted_cat = sorted(category_stats.items(), key=lambda x: x[1]['gross'], reverse=True)
    for c, v in sorted_cat:
        category_data.append({
            'name': c,
            'gross': round(v['gross'], 0),
            'qty': round(v['qty'], 0),
            'skus_count': len(v['skus']),
            'customers': len(v['customers']),
            'orders': len(v['orders']),
            'pct_of_total': round((v['gross'] / total_gross * 100) if total_gross > 0 else 0, 2)
        })

    segment_summary = defaultdict(lambda: {'count': 0, 'net': 0, 'gross': 0, 'orders': 0, 'customers': []})
    for c in sorted_customers:
        seg = c['segment']
        segment_summary[seg]['count'] += 1
        segment_summary[seg]['net'] += c['net']
        segment_summary[seg]['gross'] += c['gross']
        segment_summary[seg]['orders'] += len(c['orders'])
        segment_summary[seg]['color'] = c['segment_color']
        segment_summary[seg]['code'] = c['segment_code']
        segment_summary[seg]['action'] = c['action']
        if len(segment_summary[seg]['customers']) < 5:
            segment_summary[seg]['customers'].append(c['name'])

    segment_data = []
    for seg, v in sorted(segment_summary.items(), key=lambda x: x[1]['net'], reverse=True):
        segment_data.append({
            'name': seg,
            'code': v['code'],
            'color': v['color'],
            'action': v['action'],
            'count': v['count'],
            'net': round(v['net'], 0),
            'gross': round(v['gross'], 0),
            'orders': v['orders'],
            'cust_pct': round((v['count'] / len(customers) * 100), 1),
            'rev_pct': round((v['net'] / total_net * 100), 1),
            'sample_customers': v['customers']
        })

    abc_summary = {
        'A': {'count': 0, 'net': 0, 'label': 'Hạng A (Top 80% Doanh thu - Cốt lõi)', 'color': '#2563eb'},
        'B': {'count': 0, 'net': 0, 'label': 'Hạng B (15% Tiếp theo - Trọng tâm)', 'color': '#0ea5e9'},
        'C': {'count': 0, 'net': 0, 'label': 'Hạng C (5% Cuối cùng - Khách lẻ)', 'color': '#94a3b8'}
    }
    for c in sorted_customers:
        abc_summary[c['abc']]['count'] += 1
        abc_summary[c['abc']]['net'] += c['net']

    abc_data = []
    for tier, v in abc_summary.items():
        abc_data.append({
            'tier': tier,
            'label': v['label'],
            'color': v['color'],
            'count': v['count'],
            'net': round(v['net'], 0),
            'cust_pct': round((v['count'] / len(customers) * 100), 1),
            'rev_pct': round((v['net'] / total_net * 100), 1)
        })

    product_abc_data = [
        {
            'tier': 'A',
            'label': 'Hạng A (Top 80% Doanh thu)',
            'color': '#2563eb',
            'count': prod_abc_summary['A']['count'],
            'gross': round(prod_abc_summary['A']['gross'], 0),
            'qty': round(prod_abc_summary['A']['qty'], 0),
            'sku_pct': round((prod_abc_summary['A']['count'] / len(products) * 100), 1),
            'rev_pct': round((prod_abc_summary['A']['gross'] / total_gross * 100), 1)
        },
        {
            'tier': 'B',
            'label': 'Hạng B (15% Tiếp theo)',
            'color': '#0ea5e9',
            'count': prod_abc_summary['B']['count'],
            'gross': round(prod_abc_summary['B']['gross'], 0),
            'qty': round(prod_abc_summary['B']['qty'], 0),
            'sku_pct': round((prod_abc_summary['B']['count'] / len(products) * 100), 1),
            'rev_pct': round((prod_abc_summary['B']['gross'] / total_gross * 100), 1)
        },
        {
            'tier': 'C',
            'label': 'Hạng C (Chậm luân chuyển 5%)',
            'color': '#94a3b8',
            'count': prod_abc_summary['C']['count'],
            'gross': round(prod_abc_summary['C']['gross'], 0),
            'qty': round(prod_abc_summary['C']['qty'], 0),
            'sku_pct': round((prod_abc_summary['C']['count'] / len(products) * 100), 1),
            'rev_pct': round((prod_abc_summary['C']['gross'] / total_gross * 100), 1)
        }
    ]

    final_payload = {
        'kpi': kpi_summary,
        'monthly': monthly_data,
        'abc': abc_data,
        'segments': segment_data,
        'provinces': province_data,
        'sales_reps': sales_rep_data,
        'categories': category_data,
        'customers': formatted_customers,
        'products': formatted_products,
        'product_abc': product_abc_data,
        'predictions': reorder_predictions_sorted,
        'prediction_summary': prediction_summary
    }

    out_file = '/Users/Admin/Documents/KIng BLue/Report/customer_analytics_data.json'
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(final_payload, f, ensure_ascii=False, indent=2)

    print(f"Successfully exported data to: {out_file}")
    print(f"Total customers: {len(formatted_customers)}")
    print(f"Total products: {len(formatted_products)}")
    print(f"Total predictions: {len(reorder_predictions_sorted)}")
    print(f"File size: {os.path.getsize(out_file) / 1024:.1f} KB")

if __name__ == '__main__':
    main()
