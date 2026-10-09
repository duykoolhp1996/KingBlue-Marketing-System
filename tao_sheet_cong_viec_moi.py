# -*- coding: utf-8 -*-
"""
TẠO GOOGLE SHEET QUẢN LÝ CÔNG VIỆC - KING BLUE MARKETING
========================================================
Cấu trúc: MỖI NHÂN SỰ 1 TAB RIÊNG + 1 TAB "TONG HOP" TỰ ĐỘNG GỘP LẠI.

  · 10 tab cá nhân : Marketing Manager, Hoai Thuong, Kieu Thuong, Thuy Thuong,
                     Thien, Tu, Thuc, Ngan, Hung MN, Hung MB
  · CAU HINH       : danh mục nuôi dropdown (trạng thái, ưu tiên, nhóm việc, người phối hợp)
  · TONG HOP       : tự động gộp công việc của TẤT CẢ mọi người bằng công thức

Mỗi tab cá nhân tự động:
  - Đánh số STT
  - Điền sẵn tên người thực hiện
  - Tự tính "Trạng thái hạn" (Quá hạn / Sắp đến hạn / Đúng hạn / Đã xong)

Cách dùng:
  python3 tao_sheet_cong_viec_moi.py            # Tạo trực tiếp trên Google Drive (cần mạng)
  python3 tao_sheet_cong_viec_moi.py --xlsx     # Xuất file .xlsx để tự upload (không cần mạng)
"""

import json
import os
import sys
import datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TOKEN_FILE = os.path.join(BASE_DIR, "API", "token.json")
CONFIG_FILE = os.path.join(BASE_DIR, "cong_viec_config.json")
LINKS_FILE = os.path.join(BASE_DIR, "google_sheets_links.json")

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]

TITLE = "Quản Lý Công Việc - King Blue Marketing"
TAB_CONFIG = "CAU HINH"
TAB_TOTAL = "TONG HOP"

MAX_ROW = 500          # dòng cuối cùng vùng dữ liệu
HEADER_ROW = 4         # dòng tiêu đề cột
DATA_ROW = 5           # dòng bắt đầu dữ liệu
NCOL = 14

STATUSES = ["Chưa bắt đầu", "Đang làm", "Đang duyệt", "Hoàn thành", "Tạm hoãn"]
PRIORITIES = ["Khẩn cấp", "Cao", "Trung bình", "Thấp"]
CATEGORIES = [
    "Website & SEO", "Facebook Fanpage", "YouTube / Video", "TikTok & KOC",
    "Trade & POSM Đại lý", "Sàn TMĐT (Shopee/Lazada)", "Thiết kế & Media",
    "Đóng gói & Kho vận", "Khác",
]

# --------------------------------------------------------------------------- #
# Danh sách nhân sự -> mỗi người 1 tab
# --------------------------------------------------------------------------- #
PEOPLE = [
    {"tab": "Marketing Manager", "name": "Marketing Manager", "role": "Trưởng Phòng Marketing", "group": "Ban Quản Lý"},
    {"tab": "Hoai Thuong", "name": "Hoài Thương", "role": "Content Marketing & SEO", "group": "Content"},
    {"tab": "Kieu Thuong", "name": "Kiều Thương", "role": "Content Video & KOC", "group": "Content"},
    {"tab": "Thuy Thuong", "name": "Thụy Thương", "role": "Trade Marketing", "group": "Content"},
    {"tab": "Thien", "name": "Thiện", "role": "Graphic Designer", "group": "Design"},
    {"tab": "Tu", "name": "Tứ", "role": "Media / Photographer", "group": "Media"},
    {"tab": "Thuc", "name": "Thức", "role": "Vận hành Sàn TMĐT", "group": "Sàn TMĐT"},
    {"tab": "Ngan", "name": "Ngân", "role": "CSKH & TikTok Shop", "group": "Sàn TMĐT"},
    {"tab": "Hung MN", "name": "Hùng", "role": "Đóng gói & Kho Miền Nam", "group": "Đóng gói kho"},
    {"tab": "Hung MB", "name": "Hùng Miền Bắc", "role": "Đóng gói & Kho Miền Bắc", "group": "Đóng gói kho"},
]
SUPPORTERS = [p["name"] for p in PEOPLE]

HEADERS = [
    "STT", "Mã CV", "Tên công việc", "Nhóm công việc", "Người thực hiện",
    "Người phối hợp", "Ngày bắt đầu", "Hạn chót", "Trạng thái",
    "% Tiến độ", "Trạng thái hạn", "Độ ưu tiên", "Link tài liệu", "Ghi chú",
]
CONFIG_HEADERS = ["TRẠNG THÁI", "ĐỘ ƯU TIÊN", "NHÓM CÔNG VIỆC", "NGƯỜI PHỐI HỢP"]
CONFIG_COLUMNS = [STATUSES, PRIORITIES, CATEGORIES, SUPPORTERS]

# Dữ liệu mẫu: mỗi người 1 việc (thứ tự cột: B, C, D, F, G, H, I, J, L, M, N)
SAMPLE = {
    "Marketing Manager": [["KB-MKT-01", "Duyệt kế hoạch Marketing tháng 10 & phân bổ KPI", "Khác", "Hoài Thương", "01/10/2026", "07/10/2026", "Đang làm", "50%", "Khẩn cấp", "", "Họp Ban Giám Đốc cuối tuần"]],
    "Hoai Thuong": [["KB-MKT-02", "Rà soát & tối ưu SEO toàn bộ trang sản phẩm", "Website & SEO", "Thiện", "01/10/2026", "10/10/2026", "Đang làm", "60%", "Cao", "", "Ưu tiên nhóm sản phẩm chủ lực"]],
    "Kieu Thuong": [["KB-MKT-03", "Sản xuất 10 video review sản phẩm đăng TikTok", "TikTok & KOC", "Tứ", "28/09/2026", "08/10/2026", "Đang duyệt", "90%", "Khẩn cấp", "", "Chờ duyệt kịch bản số 7-10"]],
    "Thuy Thuong": [["KB-MKT-04", "Hoàn thiện chính sách điểm bán cho đại lý khu vực", "Trade & POSM Đại lý", "Thiện", "25/09/2026", "05/10/2026", "Hoàn thành", "100%", "Trung bình", "", "Đã gửi Ban Giám Đốc phê duyệt"]],
    "Thien": [["KB-MKT-05", "Thiết kế bộ POSM cho chiến dịch Tết", "Thiết kế & Media", "Thụy Thương", "30/09/2026", "15/10/2026", "Đang làm", "40%", "Trung bình", "", "Banner, standee, catalogue"]],
    "Tu": [["KB-MKT-06", "Chụp bộ ảnh sản phẩm máy khoan & máy cắt", "Thiết kế & Media", "Thiện", "02/10/2026", "12/10/2026", "Đang làm", "35%", "Cao", "", "Chụp tại xưởng, 3 bối cảnh"]],
    "Thuc": [["KB-MKT-07", "Đối soát tồn kho & đẩy đơn Shopee/TikTok Shop", "Sàn TMĐT (Shopee/Lazada)", "Ngân", "02/10/2026", "09/10/2026", "Chưa bắt đầu", "0%", "Cao", "", "Kiểm tra tồn kho miền Nam"]],
    "Ngan": [["KB-MKT-08", "Tối ưu tỷ lệ phản hồi chat & đánh giá 5 sao", "Sàn TMĐT (Shopee/Lazada)", "Thức", "01/10/2026", "31/10/2026", "Đang làm", "70%", "Trung bình", "", "Mục tiêu phản hồi > 95%"]],
    "Hung MN": [["KB-MKT-09", "Chuẩn hoá quy cách đóng gói chống sốc đơn online", "Đóng gói & Kho vận", "", "29/09/2026", "06/10/2026", "Đang làm", "80%", "Cao", "", "Kho tổng miền Nam"]],
    "Hung MB": [["KB-MKT-10", "Kiểm kê vật tư đóng gói kho Hà Nội", "Đóng gói & Kho vận", "", "30/09/2026", "14/10/2026", "Chưa bắt đầu", "0%", "Thấp", "", "Bao bì, băng keo, xốp hơi"]],
}


# --------------------------------------------------------------------------- #
# Sinh công thức
# --------------------------------------------------------------------------- #
def f_stt(sheets_mode):
    expr = f'IF(B{DATA_ROW}:B{MAX_ROW}<>"",ROW(B{DATA_ROW}:B{MAX_ROW})-{DATA_ROW-1},"")'
    return f"=ARRAYFORMULA({expr})" if sheets_mode else f"={expr}"


def f_assignee(name, sheets_mode):
    expr = f'IF(B{DATA_ROW}:B{MAX_ROW}<>"","{name}","")'
    return f"=ARRAYFORMULA({expr})" if sheets_mode else f"={expr}"


def f_deadline(sheets_mode):
    expr = (
        f'IF(B{DATA_ROW}:B{MAX_ROW}="","",'
        f'IF(I{DATA_ROW}:I{MAX_ROW}="Hoàn thành","Đã xong",'
        f'IFERROR(IF(H{DATA_ROW}:H{MAX_ROW}="","Đúng hạn",'
        f'IF(H{DATA_ROW}:H{MAX_ROW}<TODAY(),"Quá hạn",'
        f'IF(H{DATA_ROW}:H{MAX_ROW}-TODAY()<=2,"Sắp đến hạn","Đúng hạn"))),"Đúng hạn")))'
    )
    return f"=ARRAYFORMULA({expr})" if sheets_mode else f"={expr}"


def f_total(sheets_mode):
    """Công thức gộp toàn bộ tab cá nhân vào tab TONG HOP."""
    if sheets_mode:
        parts = "; ".join(
            f"'{p['tab']}'!A{DATA_ROW}:N{MAX_ROW}" for p in PEOPLE
        )
        return f'=QUERY({{{parts}}}, "select * where Col2 is not null", 0)'
    stack_all = "VSTACK(" + ", ".join(f"'{p['tab']}'!A{DATA_ROW}:N{MAX_ROW}" for p in PEOPLE) + ")"
    stack_key = "VSTACK(" + ", ".join(f"'{p['tab']}'!B{DATA_ROW}:B{MAX_ROW}" for p in PEOPLE) + ")"
    return f'=FILTER({stack_all}, ({stack_key}<>"")*({stack_key}<>0))'


def f_stat(field, value, agg="SUM"):
    parts = ", ".join(f'COUNTIF(\'{p["tab"]}\'!{field}{DATA_ROW}:{field}{MAX_ROW},"{value}")' for p in PEOPLE)
    return f"={agg}({parts})"


def f_count_nonblank():
    parts = ", ".join(f'COUNTA(\'{p["tab"]}\'!B{DATA_ROW}:B{MAX_ROW})' for p in PEOPLE)
    return f"=SUM({parts})"


def col_letter(n):
    s = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


LAST_COL = col_letter(NCOL)


# --------------------------------------------------------------------------- #
# Xuất file .xlsx (không cần mạng)
# --------------------------------------------------------------------------- #
def build_xlsx(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    title_font = Font(bold=True, color="FFFFFF", size=13)
    head_font = Font(bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="DCE1E9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    band = PatternFill("solid", fgColor="F6F8FC")
    widths = [6, 12, 42, 22, 17, 16, 14, 14, 15, 11, 15, 13, 22, 30]

    # ---------------- Tab cá nhân ----------------
    for p in PEOPLE:
        ws = wb.create_sheet(p["tab"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)
        c = ws.cell(row=1, column=1, value=f"CÔNG VIỆC - {p['name'].upper()}  ({p['role']})")
        c.font = title_font
        c.fill = PatternFill("solid", fgColor="0A1329")
        c.alignment = center
        ws.row_dimensions[1].height = 34
        ws.cell(row=2, column=1, value=f"Bộ phận: {p['group']}  ·  Mọi thay đổi ở đây tự động tổng hợp vào tab TONG HOP.").font = \
            Font(italic=True, color="64748B", size=9)

        for i, h in enumerate(HEADERS, start=1):
            cell = ws.cell(row=HEADER_ROW, column=i, value=h)
            cell.font = head_font
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
            cell.alignment = center
            cell.border = border
        ws.row_dimensions[HEADER_ROW].height = 30

        for r, row in enumerate(SAMPLE.get(p["tab"], []), start=DATA_ROW):
            # Ánh xạ 11 giá trị mẫu -> 14 cột (A, E, K để trống cho công thức tự động)
            vals = [None, row[0], row[1], row[2], None,
                    row[3], row[4], row[5], row[6], row[7], None,
                    row[8], row[9], row[10]]
            for i, val in enumerate(vals, start=1):
                if i in (1, 5, 11):
                    continue
                if i in (7, 8) and val:
                    d, m, y = str(val).split("/")
                    val = datetime.date(int(y), int(m), int(d))
                cell = ws.cell(row=r, column=i, value=val)
                cell.border = border
                cell.alignment = wrap if i in (3, 14) else Alignment(vertical="center")
                if i in (7, 8):
                    cell.number_format = "DD/MM/YYYY"

        # Kẻ bảng + tô xen kẽ cho khu vực nhập liệu
        for r in range(DATA_ROW, 81):
            for i in range(1, NCOL + 1):
                cell = ws.cell(row=r, column=i)
                cell.border = border
                if (r - DATA_ROW) % 2 == 1:
                    cell.fill = band

        ws.cell(row=DATA_ROW, column=1, value=f_stt(False))
        ws.cell(row=DATA_ROW, column=5, value=f_assignee(p["name"], False))
        ws.cell(row=DATA_ROW, column=11, value=f_deadline(False))

        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{DATA_ROW}"

        def dv(col, source, title):
            d = DataValidation(type="list", formula1=source, allow_blank=True)
            d.error = "Giá trị không hợp lệ, vui lòng chọn từ danh sách."
            d.promptTitle = title
            ws.add_data_validation(d)
            d.add(f"{col}{DATA_ROW}:{col}{MAX_ROW}")

        dv("D", f"='{TAB_CONFIG}'!$C$4:$C$40", "Nhóm công việc")
        dv("F", f"='{TAB_CONFIG}'!$D$4:$D$40", "Người phối hợp")
        dv("I", f"='{TAB_CONFIG}'!$A$4:$A$40", "Trạng thái")
        dv("L", f"='{TAB_CONFIG}'!$B$4:$B$40", "Độ ưu tiên")

    # ---------------- Tab CAU HINH ----------------
    ws2 = wb.create_sheet(TAB_CONFIG)
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    t = ws2.cell(row=1, column=1, value="DANH MỤC CẤU HÌNH - KING BLUE MARKETING")
    t.font = Font(bold=True, color="FFFFFF", size=12)
    t.fill = PatternFill("solid", fgColor="0F172A")
    t.alignment = center
    for i, h in enumerate(CONFIG_HEADERS, start=1):
        cell = ws2.cell(row=3, column=i, value=h)
        cell.font = head_font
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = center
        cell.border = border
    for r in range(max(len(c) for c in CONFIG_COLUMNS)):
        for i, col in enumerate(CONFIG_COLUMNS, start=1):
            cell = ws2.cell(row=4 + r, column=i, value=col[r] if r < len(col) else None)
            cell.border = border
    for i in range(1, 5):
        ws2.column_dimensions[get_column_letter(i)].width = 26

    # ---------------- Tab TONG HOP ----------------
    ws3 = wb.create_sheet(TAB_TOTAL)
    ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)
    d = ws3.cell(row=1, column=1, value="TỔNG HỢP CÔNG VIỆC TOÀN PHÒNG MARKETING - KING BLUE")
    d.font = title_font
    d.fill = PatternFill("solid", fgColor="0A1329")
    d.alignment = center
    ws3.row_dimensions[1].height = 34

    stats = [
        ("TỔNG SỐ VIỆC", f_count_nonblank()),
        ("HOÀN THÀNH", f_stat("I", "Hoàn thành")),
        ("ĐANG LÀM", f_stat("I", "Đang làm")),
        ("ĐANG DUYỆT", f_stat("I", "Đang duyệt")),
        ("QUÁ HẠN", f_stat("K", "Quá hạn")),
        ("SẮP ĐẾN HẠN", f_stat("K", "Sắp đến hạn")),
    ]
    for i, (label, _) in enumerate(stats, start=1):
        cell = ws3.cell(row=3, column=i, value=label)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = center
    for i, (_, formula) in enumerate(stats, start=1):
        cell = ws3.cell(row=4, column=i, value=formula)
        cell.font = Font(bold=True, size=14, color="1D4ED8")
        cell.alignment = center

    for i, h in enumerate(HEADERS, start=1):
        cell = ws3.cell(row=6, column=i, value=h)
        cell.font = head_font
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = center
        cell.border = border
    for r in range(7, 81):
        for i in range(1, NCOL + 1):
            cell = ws3.cell(row=r, column=i)
            cell.border = border
            if (r - 7) % 2 == 1:
                cell.fill = band

    ws3.cell(row=7, column=1, value=f_total(False))
    for i, w in enumerate(widths, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A7"

    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# Tạo trực tiếp trên Google Drive
# --------------------------------------------------------------------------- #
def authenticate():
    from google.oauth2.credentials import Credentials
    from google.auth.transport.requests import Request
    if not os.path.exists(TOKEN_FILE):
        sys.exit(f"❌ Không tìm thấy {TOKEN_FILE}. Hãy chạy connect_google_api.py trước.")
    creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    if not creds.valid:
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())
            with open(TOKEN_FILE, "w", encoding="utf-8") as f:
                f.write(creds.to_json())
            print("🔄 Đã làm mới token Google API.")
        else:
            sys.exit("❌ Token không hợp lệ. Hãy chạy lại connect_google_api.py.")
    return creds


def create_spreadsheet(sheets):
    titles = [p["tab"] for p in PEOPLE] + [TAB_CONFIG, TAB_TOTAL]
    body = {
        "properties": {"title": TITLE, "locale": "vi_VN", "timeZone": "Asia/Ho_Chi_Minh"},
        "sheets": [{"properties": {"title": t, "gridProperties": {"rowCount": 400, "columnCount": NCOL}}}
                   for t in titles],
    }
    ss = sheets.spreadsheets().create(body=body, fields="spreadsheetId,properties.title").execute()
    print(f"✅ Đã tạo Spreadsheet: {ss['properties']['title']}")
    print(f"   ID: {ss['spreadsheetId']}")
    return ss["spreadsheetId"]


def write_data(sheets, sid):
    # Tab cá nhân
    for p in PEOPLE:
        sheets.spreadsheets().values().update(
            spreadsheetId=sid, range=f"'{p['tab']}'!A1",
            valueInputOption="USER_ENTERED",
            body={"values": [[f"CÔNG VIỆC - {p['name'].upper()}  ({p['role']})"]]},
        ).execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sid, range=f"'{p['tab']}'!A{HEADER_ROW}:{LAST_COL}{HEADER_ROW}",
            valueInputOption="USER_ENTERED", body={"values": [HEADERS]},
        ).execute()
        rows = SAMPLE.get(p["tab"], [])
        if rows:
            block_bd, block_fn = [], []
            for row in rows:
                block_bd.append(row[:3])
                block_fn.append(row[3:])
            sheets.spreadsheets().values().update(
                spreadsheetId=sid, range=f"'{p['tab']}'!B{DATA_ROW}:D{DATA_ROW + len(rows) - 1}",
                valueInputOption="USER_ENTERED", body={"values": block_bd},
            ).execute()
            sheets.spreadsheets().values().update(
                spreadsheetId=sid, range=f"'{p['tab']}'!F{DATA_ROW}:N{DATA_ROW + len(rows) - 1}",
                valueInputOption="USER_ENTERED", body={"values": block_fn},
            ).execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sid, range=f"'{p['tab']}'!A{DATA_ROW}",
            valueInputOption="USER_ENTERED", body={"values": [[f_stt(True)]]},
        ).execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sid, range=f"'{p['tab']}'!E{DATA_ROW}",
            valueInputOption="USER_ENTERED", body={"values": [[f_assignee(p["name"], True)]]},
        ).execute()
        sheets.spreadsheets().values().update(
            spreadsheetId=sid, range=f"'{p['tab']}'!K{DATA_ROW}",
            valueInputOption="USER_ENTERED", body={"values": [[f_deadline(True)]]},
        ).execute()

    # Tab CAU HINH
    cfg = [["DANH MỤC CẤU HÌNH - KING BLUE MARKETING", "", "", ""], ["", "", "", ""], CONFIG_HEADERS]
    for i in range(max(len(c) for c in CONFIG_COLUMNS)):
        cfg.append([col[i] if i < len(col) else "" for col in CONFIG_COLUMNS])
    sheets.spreadsheets().values().update(
        spreadsheetId=sid, range=f"'{TAB_CONFIG}'!A1:D{len(cfg)}",
        valueInputOption="USER_ENTERED", body={"values": cfg},
    ).execute()

    # Tab TONG HOP
    stats = [
        ("TỔNG SỐ VIỆC", f_count_nonblank()), ("HOÀN THÀNH", f_stat("I", "Hoàn thành")),
        ("ĐANG LÀM", f_stat("I", "Đang làm")), ("ĐANG DUYỆT", f_stat("I", "Đang duyệt")),
        ("QUÁ HẠN", f_stat("K", "Quá hạn")), ("SẮP ĐẾN HẠN", f_stat("K", "Sắp đến hạn")),
    ]
    sheets.spreadsheets().values().update(
        spreadsheetId=sid, range=f"'{TAB_TOTAL}'!A1",
        valueInputOption="USER_ENTERED",
        body={"values": [["TỔNG HỢP CÔNG VIỆC TOÀN PHÒNG MARKETING - KING BLUE"]]},
    ).execute()
    sheets.spreadsheets().values().update(
        spreadsheetId=sid, range=f"'{TAB_TOTAL}'!A3:F3",
        valueInputOption="USER_ENTERED", body={"values": [[s[0] for s in stats]]},
    ).execute()
    sheets.spreadsheets().values().update(
        spreadsheetId=sid, range=f"'{TAB_TOTAL}'!A4:F4",
        valueInputOption="USER_ENTERED", body={"values": [[s[1] for s in stats]]},
    ).execute()
    sheets.spreadsheets().values().update(
        spreadsheetId=sid, range=f"'{TAB_TOTAL}'!A6:N6",
        valueInputOption="USER_ENTERED", body={"values": [HEADERS]},
    ).execute()
    sheets.spreadsheets().values().update(
        spreadsheetId=sid, range=f"'{TAB_TOTAL}'!A7",
        valueInputOption="USER_ENTERED", body={"values": [[f_total(True)]]},
    ).execute()
    print("✅ Đã ghi dữ liệu, công thức tự động và tab tổng hợp.")


def format_sheets(sheets, sid):
    meta = sheets.spreadsheets().get(spreadsheetId=sid).execute()
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}
    reqs = []
    widths = [46, 100, 320, 180, 140, 130, 110, 110, 120, 90, 120, 110, 160, 240]
    title_bg = {"red": 0.039, "green": 0.075, "blue": 0.160}
    head_bg = {"red": 0.117, "green": 0.227, "blue": 0.541}
    line = {"red": 0.86, "green": 0.88, "blue": 0.91}

    for p in PEOPLE:
        sid_ = ids[p["tab"]]
        reqs += [
            {"mergeCells": {"range": {"sheetId": sid_, "startRowIndex": 0, "endRowIndex": 1,
                                      "startColumnIndex": 0, "endColumnIndex": NCOL}, "mergeType": "MERGE_ALL"}},
            {"repeatCell": {"range": {"sheetId": sid_, "startRowIndex": 0, "endRowIndex": 1,
                                      "startColumnIndex": 0, "endColumnIndex": NCOL},
                            "cell": {"userEnteredFormat": {"backgroundColor": title_bg, "horizontalAlignment": "CENTER",
                                     "verticalAlignment": "MIDDLE", "textFormat": {"bold": True, "fontSize": 13,
                                     "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
            {"updateDimensionProperties": {"range": {"sheetId": sid_, "dimension": "ROWS", "startIndex": 0, "endIndex": 1},
                                           "properties": {"pixelSize": 42}, "fields": "pixelSize"}},
            {"repeatCell": {"range": {"sheetId": sid_, "startRowIndex": 3, "endRowIndex": 4,
                                      "startColumnIndex": 0, "endColumnIndex": NCOL},
                            "cell": {"userEnteredFormat": {"backgroundColor": head_bg, "horizontalAlignment": "CENTER",
                                     "verticalAlignment": "MIDDLE", "textFormat": {"bold": True, "fontSize": 10,
                                     "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
            {"updateDimensionProperties": {"range": {"sheetId": sid_, "dimension": "ROWS", "startIndex": 3, "endIndex": 4},
                                           "properties": {"pixelSize": 36}, "fields": "pixelSize"}},
            {"updateSheetProperties": {"properties": {"sheetId": sid_, "gridProperties": {"frozenRowCount": HEADER_ROW}},
                                       "fields": "gridProperties.frozenRowCount"}},
            {"repeatCell": {"range": {"sheetId": sid_, "startRowIndex": 4, "endRowIndex": 400,
                                      "startColumnIndex": 0, "endColumnIndex": NCOL},
                            "cell": {"userEnteredFormat": {"borders": {
                                "top": {"style": "SOLID", "color": line}, "bottom": {"style": "SOLID", "color": line},
                                "left": {"style": "SOLID", "color": line}, "right": {"style": "SOLID", "color": line}}}},
                            "fields": "userEnteredFormat.borders"}},
            {"addBanding": {"bandedRange": {"range": {"sheetId": sid_, "startRowIndex": 3, "endRowIndex": 400,
                                                      "startColumnIndex": 0, "endColumnIndex": NCOL},
                                            "rowProperties": {"firstBandColor": {"red": 1, "green": 1, "blue": 1},
                                                              "secondBandColor": {"red": 0.965, "green": 0.973, "blue": 0.988}}}}},
            {"repeatCell": {"range": {"sheetId": sid_, "startRowIndex": 4, "endRowIndex": 400,
                                      "startColumnIndex": 2, "endColumnIndex": 3},
                            "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "MIDDLE"}},
                            "fields": "userEnteredFormat.wrapStrategy,userEnteredFormat.verticalAlignment"}},
            {"repeatCell": {"range": {"sheetId": sid_, "startRowIndex": 4, "endRowIndex": 400,
                                      "startColumnIndex": 6, "endColumnIndex": 8},
                            "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "dd/mm/yyyy"}}},
                            "fields": "userEnteredFormat.numberFormat"}},
        ]
        for i, w in enumerate(widths):
            reqs.append({"updateDimensionProperties": {
                "range": {"sheetId": sid_, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
                "properties": {"pixelSize": w}, "fields": "pixelSize"}})
        for col, src in [("D", f"='{TAB_CONFIG}'!$C$4:$C$40"), ("F", f"='{TAB_CONFIG}'!$D$4:$D$40"),
                         ("I", f"='{TAB_CONFIG}'!$A$4:$A$40"), ("L", f"='{TAB_CONFIG}'!$B$4:$B$40")]:
            idx = ord(col) - 65
            reqs.append({"setDataValidation": {
                "range": {"sheetId": sid_, "startRowIndex": 4, "endRowIndex": 400,
                          "startColumnIndex": idx, "endColumnIndex": idx + 1},
                "rule": {"condition": {"type": "ONE_OF_RANGE", "values": [{"userEnteredValue": src}]},
                         "showCustomUi": True, "strict": False}}})

    # CAU HINH
    cid = ids[TAB_CONFIG]
    reqs += [
        {"mergeCells": {"range": {"sheetId": cid, "startRowIndex": 0, "endRowIndex": 1,
                                  "startColumnIndex": 0, "endColumnIndex": 4}, "mergeType": "MERGE_ALL"}},
        {"repeatCell": {"range": {"sheetId": cid, "startRowIndex": 0, "endRowIndex": 1,
                                  "startColumnIndex": 0, "endColumnIndex": 4},
                        "cell": {"userEnteredFormat": {"backgroundColor": {"red": 0.058, "green": 0.090, "blue": 0.165},
                                 "horizontalAlignment": "CENTER", "verticalAlignment": "MIDDLE",
                                 "textFormat": {"bold": True, "fontSize": 12,
                                 "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
        {"repeatCell": {"range": {"sheetId": cid, "startRowIndex": 2, "endRowIndex": 3,
                                  "startColumnIndex": 0, "endColumnIndex": 4},
                        "cell": {"userEnteredFormat": {"backgroundColor": head_bg, "horizontalAlignment": "CENTER",
                                 "verticalAlignment": "MIDDLE", "textFormat": {"bold": True, "fontSize": 10,
                                 "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
    ]
    for i in range(4):
        reqs.append({"updateDimensionProperties": {
            "range": {"sheetId": cid, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
            "properties": {"pixelSize": 200}, "fields": "pixelSize"}})

    # TONG HOP
    tid = ids[TAB_TOTAL]
    reqs += [
        {"mergeCells": {"range": {"sheetId": tid, "startRowIndex": 0, "endRowIndex": 1,
                                  "startColumnIndex": 0, "endColumnIndex": NCOL}, "mergeType": "MERGE_ALL"}},
        {"repeatCell": {"range": {"sheetId": tid, "startRowIndex": 0, "endRowIndex": 1,
                                  "startColumnIndex": 0, "endColumnIndex": NCOL},
                        "cell": {"userEnteredFormat": {"backgroundColor": title_bg, "horizontalAlignment": "CENTER",
                                 "verticalAlignment": "MIDDLE", "textFormat": {"bold": True, "fontSize": 13,
                                 "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
        {"repeatCell": {"range": {"sheetId": tid, "startRowIndex": 2, "endRowIndex": 3,
                                  "startColumnIndex": 0, "endColumnIndex": 6},
                        "cell": {"userEnteredFormat": {"backgroundColor": head_bg, "horizontalAlignment": "CENTER",
                                 "verticalAlignment": "MIDDLE", "textFormat": {"bold": True, "fontSize": 9,
                                 "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
        {"repeatCell": {"range": {"sheetId": tid, "startRowIndex": 3, "endRowIndex": 4,
                                  "startColumnIndex": 0, "endColumnIndex": 6},
                        "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER", "verticalAlignment": "MIDDLE",
                                 "textFormat": {"bold": True, "fontSize": 14,
                                 "foregroundColor": {"red": 0.114, "green": 0.306, "blue": 0.847}}}},
                        "fields": "userEnteredFormat"}},
        {"repeatCell": {"range": {"sheetId": tid, "startRowIndex": 5, "endRowIndex": 6,
                                  "startColumnIndex": 0, "endColumnIndex": NCOL},
                        "cell": {"userEnteredFormat": {"backgroundColor": head_bg, "horizontalAlignment": "CENTER",
                                 "verticalAlignment": "MIDDLE", "textFormat": {"bold": True, "fontSize": 10,
                                 "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}}, "fields": "userEnteredFormat"}},
        {"updateSheetProperties": {"properties": {"sheetId": tid, "gridProperties": {"frozenRowCount": 6}},
                                   "fields": "gridProperties.frozenRowCount"}},
    ]
    for i, w in enumerate(widths):
        reqs.append({"updateDimensionProperties": {
            "range": {"sheetId": tid, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
            "properties": {"pixelSize": w}, "fields": "pixelSize"}})

    sheets.spreadsheets().batchUpdate(spreadsheetId=sid, body={"requests": reqs}).execute()
    print(f"✅ Đã định dạng {len(reqs)} thuộc tính (màu, kẻ dòng, dropdown, ngày tháng).")


def move_to_folder(creds, sid):
    folder_id = None
    if os.path.exists(LINKS_FILE):
        try:
            with open(LINKS_FILE, "r", encoding="utf-8") as f:
                folder_id = json.load(f).get("folder_id")
        except Exception:
            folder_id = None
    if not folder_id:
        return None
    try:
        from googleapiclient.discovery import build
        drive = build("drive", "v3", credentials=creds)
        prev = ",".join(drive.files().get(fileId=sid, fields="parents").execute().get("parents", []))
        drive.files().update(fileId=sid, addParents=folder_id, removeParents=prev, fields="id,parents").execute()
        print(f"✅ Đã di chuyển file vào thư mục Drive King Blue.")
    except Exception as e:
        print(f"⚠️  Không di chuyển được vào thư mục (vẫn dùng bình thường): {e}")
    return folder_id


def save_config(sid, folder_id):
    cfg = {
        "spreadsheet_id": sid,
        "spreadsheet_title": TITLE,
        "spreadsheet_url": f"https://docs.google.com/spreadsheets/d/{sid}/edit",
        "folder_id": folder_id,
        "tabs": {"people": [p["tab"] for p in PEOPLE], "config": TAB_CONFIG, "total": TAB_TOTAL},
        "header_row": HEADER_ROW,
        "data_start_row": DATA_ROW,
        "created_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    print(f"\n🔗 LINK GOOGLE SHEET: {cfg['spreadsheet_url']}\n")
    return cfg


def main():
    args = [a.lower() for a in sys.argv[1:]]
    print("=" * 74)
    print("  GOOGLE SHEET QUẢN LÝ CÔNG VIỆC - KING BLUE MARKETING")
    print(f"  Cấu trúc: {len(PEOPLE)} tab nhân sự  +  1 tab '{TAB_TOTAL}' tự động gộp")
    print("=" * 74)

    if "--xlsx" in args or "xlsx" in args:
        out = os.path.join(BASE_DIR, "Quan_Ly_Cong_Viec_KingBlue_MOI.xlsx")
        build_xlsx(out)
        print(f"✅ Đã xuất file Excel: {out}")
        print(f"   Gồm {len(PEOPLE)} tab nhân sự + tab '{TAB_TOTAL}' + tab '{TAB_CONFIG}'")
        print("\n📌 Đưa lên Google Sheet (không cần API):")
        print("   1) Mở https://drive.google.com  ->  kéo thả file .xlsx")
        print("   2) Chuột phải file  ->  Mở bằng  ->  Google Trang tính")
        print("   3) Tệp  ->  Lưu thành Google Trang tính")
        return

    from googleapiclient.discovery import build
    creds = authenticate()
    sheets = build("sheets", "v4", credentials=creds)
    sid = create_spreadsheet(sheets)
    write_data(sheets, sid)
    format_sheets(sheets, sid)
    folder_id = move_to_folder(creds, sid)
    save_config(sid, folder_id)
    print("🎉 HOÀN TẤT! Mở link ở trên để sử dụng.")


if __name__ == "__main__":
    main()
